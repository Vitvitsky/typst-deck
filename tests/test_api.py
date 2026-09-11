import asyncio
import io
import json
from urllib.parse import quote

import httpx
import pytest
from pptx import Presentation

from tests.conftest import requires_typst
from typst_deck.llm.client import Completion, ToolCall
from typst_deck.model import Deck


async def wait_job(client, job_id):
    for _ in range(300):
        j = (await client.get(f"/api/jobs/{job_id}")).json()
        if j["status"] in ("done", "failed"):
            return j
        await asyncio.sleep(0.02)
    raise AssertionError("задача не завершилась")


async def test_upload_and_errors(client):
    r = await client.post("/api/uploads", files={"file": ("brief.txt", b"hello", "text/plain")})
    assert r.status_code == 200 and r.json()["kind"] == "text" and r.json()["chars"] == 5
    r = await client.post("/api/uploads", files={"file": ("x.exe", b"MZ", "application/octet-stream")})
    assert r.status_code == 422 and r.json()["error"]["code"] == "validation"
    assert (await client.get("/api/decks/nope")).status_code == 404
    r = await client.post("/api/decks", json={"brief": 5})
    assert r.status_code == 422 and r.json()["error"]["code"] == "validation"


async def test_create_without_llm_is_503(settings):
    from typst_deck.app import create_app

    app = create_app(settings, llm=None)
    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://t") as c:
        r = await c.post("/api/decks", json={"brief": "x", "upload_ids": []})
        assert r.status_code == 503 and r.json()["error"]["code"] == "llm_unavailable"
        h = (await c.get("/api/health")).json()
        assert h["ok"] is True and h["llm"] is False and h["model"] is None


@requires_typst
async def test_full_flow(client, fake_llm, png_bytes):
    fake_llm._queue += [
        Completion("", [ToolCall("c1", "set_deck", {"title": "Итоги", "slides": [
            {"type": "title", "title": "Итоги", "subtitle": "Q3"},
            {"type": "image", "title": "Рост", "bullets": ["a", "b"]},
            {"type": "final"}]})]),
        Completion("", [ToolCall("c2", "finish", {"summary": "Готово: 3 слайда"})]),
    ]
    up = (await client.post("/api/uploads", files={"file": ("brief.txt", "данные".encode(), "text/plain")})).json()
    r = await client.post("/api/decks", json={"brief": "сделай", "upload_ids": [up["upload_id"]]})
    assert r.status_code == 202
    deck_id, job_id = r.json()["deck_id"], r.json()["job_id"]
    assert (await client.get(f"/api/decks/{deck_id}")).json()["status"] in ("generating", "ready")
    job = await wait_job(client, job_id)
    assert job["status"] == "done", job
    d = (await client.get(f"/api/decks/{deck_id}")).json()
    assert d["status"] == "ready" and d["slide_count"] == 3 and d["edit_url"] == f"http://t/d/{deck_id}"
    assert "данные" in fake_llm.calls[0][-1]["content"]
    chat = (await client.get(f"/api/decks/{deck_id}/chat")).json()
    assert [m["role"] for m in chat] == ["user", "assistant"] and chat[1]["content"] == "Готово: 3 слайда"

    svg = await client.get(f"/api/decks/{deck_id}/slides/2.svg")
    assert svg.headers["content-type"].startswith("image/svg+xml")
    assert (await client.get(f"/api/decks/{deck_id}/slides/9.svg")).status_code == 404
    pdf = await client.get(f"/api/decks/{deck_id}/deck.pdf")
    assert pdf.status_code == 200 and pdf.content[:4] == b"%PDF"
    pos = (await client.get(f"/api/decks/{deck_id}/positions")).json()
    assert any(p["field"] == "bullets" for p in pos)

    sid = d["deck"]["slides"][1]["id"]
    r = await client.patch(f"/api/decks/{deck_id}/slides/{sid}", json={"patch": {"title": "Новый"}})
    assert r.status_code == 200 and r.json()["deck"]["slides"][1]["title"] == "Новый"
    assert r.json()["version"] == 3
    r = await client.patch(f"/api/decks/{deck_id}/slides/{sid}", json={"patch": {"title": "x" * 500}})
    assert r.status_code == 422
    assert (await client.patch(f"/api/decks/{deck_id}/slides/s_nope", json={"patch": {}})).status_code == 404

    a = await client.post(f"/api/decks/{deck_id}/assets", files={"file": ("p.png", png_bytes, "image/png")})
    aid = a.json()["asset_id"]
    r = await client.patch(f"/api/decks/{deck_id}/slides/{sid}", json={"patch": {"image": aid}})
    assert r.json()["warnings"] == []

    r = await client.post(f"/api/decks/{deck_id}/slides",
                          json={"slide": {"type": "section", "title": "Р"}, "after": sid})
    assert r.json()["slide_count"] == 4 and r.json()["deck"]["slides"][2]["type"] == "section"
    new_sid = r.json()["deck"]["slides"][2]["id"]
    r = await client.post(f"/api/decks/{deck_id}/slides/{new_sid}/move", json={"to_index": 0})
    assert r.json()["deck"]["slides"][0]["id"] == new_sid
    r = await client.delete(f"/api/decks/{deck_id}/slides/{new_sid}")
    assert r.json()["slide_count"] == 3

    versions = (await client.get(f"/api/decks/{deck_id}/versions")).json()
    assert [v["version"] for v in versions][:2] == [1, 2]
    r = await client.post(f"/api/decks/{deck_id}/revert", json={"version": 2})
    assert r.json()["deck"]["slides"][1]["title"] == "Рост"
    assert (await client.post(f"/api/decks/{deck_id}/revert", json={"version": 99})).status_code == 404

    fake_llm._queue += [
        Completion("", [ToolCall("c3", "update_slide", {"slide_id": sid, "patch": {"title": "Через чат"}})]),
        Completion("Поменял", []),
    ]
    gate = asyncio.Event()

    class Gated:  # задерживает ответ LLM, чтобы задача была активна во время второго запроса
        async def complete(self, messages, tools):
            await gate.wait()
            return await fake_llm.complete(messages, tools)

    client.app.state.service.llm = Gated()
    r = await client.post(f"/api/decks/{deck_id}/messages", json={"text": "поменяй заголовок"})
    assert r.status_code == 202
    r2 = await client.post(f"/api/decks/{deck_id}/messages", json={"text": "ещё"})
    assert r2.status_code == 409
    gate.set()
    job = await wait_job(client, r.json()["job_id"])
    assert job["status"] == "done", job
    d = (await client.get(f"/api/decks/{deck_id}")).json()
    assert d["deck"]["slides"][1]["title"] == "Через чат"
    assert sid in fake_llm.calls[-1][0]["content"]
    users = [m["content"] for m in fake_llm.calls[-1] if m["role"] == "user"]
    assert any("данные" in u for u in users)  # файлы видны и при правке


@requires_typst
async def test_failed_generation_reports(client, fake_llm):
    r = await client.post("/api/decks", json={"brief": "x", "upload_ids": []})
    job = await wait_job(client, r.json()["job_id"])
    assert job["status"] == "failed" and "FakeLLM" in job["error"]
    d = (await client.get(f"/api/decks/{r.json()['deck_id']}")).json()
    assert d["status"] == "failed" and d["last_error"]
    chat = (await client.get(f"/api/decks/{r.json()['deck_id']}/chat")).json()
    assert chat[-1]["role"] == "error"


async def test_list_decks(client):
    assert (await client.get("/api/decks")).json() == []
    from typst_deck.model import Deck
    store = client.app.state.store
    store.create(Deck(title="Первая", slides=[{"type": "final"}]))
    store.create(Deck(title="Вторая"))
    items = (await client.get("/api/decks")).json()
    assert [i["title"] for i in items] == ["Вторая", "Первая"]
    assert items[1]["slide_count"] == 1 and items[1]["status"] == "ready" and items[1]["has_build"] is False


@requires_typst
async def test_change_template(client):
    from typst_deck.model import Deck
    store = client.app.state.store
    d = store.create(Deck(title="Старое", slides=[{"type": "final"}]))
    r = await client.post(f"/api/decks/{d.id}/template", json={"template": "dark"})
    assert r.status_code == 200 and r.json()["deck"]["template"] == "dark"
    assert r.json()["template"]["name"] == "Тёмный" and r.json()["version"] == 2 and r.json()["has_build"]
    assert (await client.post(f"/api/decks/{d.id}/template", json={"template": "nope"})).status_code == 422
    assert (await client.get(f"/api/decks/{d.id}")).json()["deck"]["template"] == "dark"


@requires_typst
async def test_rename_materials_delete(client, fake_llm, png_bytes):
    from typst_deck.model import Deck
    store = client.app.state.store
    d = store.create(Deck(title="Старое", slides=[{"type": "final"}]))
    r = await client.patch(f"/api/decks/{d.id}", json={"title": "Новое"})
    assert r.status_code == 200 and r.json()["deck"]["title"] == "Новое" and r.json()["version"] == 2
    r = await client.post(f"/api/decks/{d.id}/uploads", files={"file": ("cat.png", png_bytes, "image/png")})
    assert r.json()["kind"] == "image" and r.json()["asset_id"].endswith(".png")
    r = await client.post(f"/api/decks/{d.id}/uploads", files={"file": ("notes.txt", "цифры".encode(), "text/plain")})
    assert r.json() == {"name": "notes.txt", "kind": "text", "chars": 5}
    assert store.list_uploads(d.id) == [("notes.txt", "цифры")]
    assert (await client.post(f"/api/decks/{d.id}/uploads", files={"file": ("x.exe", b"MZ", "x")})).status_code == 422
    assert (await client.delete(f"/api/decks/{d.id}")).status_code == 204
    assert (await client.get(f"/api/decks/{d.id}")).status_code == 404
    assert (await client.delete(f"/api/decks/{d.id}")).status_code == 404


@requires_typst
async def test_job_phase_and_changes_in_chat(client, fake_llm):
    fake_llm._queue += [
        Completion("", [ToolCall("c1", "set_deck", {"title": "T", "slides": [{"type": "title", "title": "A"}]})]),
        Completion("", [ToolCall("c2", "finish", {"summary": "Готово"})]),
    ]
    r = await client.post("/api/decks", json={"brief": "x", "upload_ids": []})
    job = await wait_job(client, r.json()["job_id"])
    assert job["status"] == "done" and job["phase"] == "render" and job["created_at"]
    chat = (await client.get(f"/api/decks/{r.json()['deck_id']}/chat")).json()
    assert chat[-1]["changes"][0]["kind"] == "rebuilt" and chat[-1]["changes"][0]["count"] == 1
    d = (await client.get(f"/api/decks/{r.json()['deck_id']}")).json()
    assert d["job"] is None


@requires_typst
async def test_templates_history_and_template_choice(client, fake_llm):
    ts = (await client.get("/api/templates")).json()
    ids = [t["id"] for t in ts]
    # только шаблоны из git: корпоративные подключаются монтированием, их проверяют
    # tests/test_corporate.py и сервис при старте
    assert ids[0] == "default" and "dark" in ids and ts[ids.index("dark")]["name"] == "Тёмный"
    svg = await client.get("/api/templates/dark/preview/1.svg")
    assert svg.status_code == 200 and svg.headers["content-type"].startswith("image/svg+xml")
    assert (await client.get("/api/templates/dark/preview/9.svg")).status_code == 404
    assert (await client.get("/api/templates/nope/preview/1.svg")).status_code == 422

    fake_llm._queue += [
        Completion("", [ToolCall("c1", "set_deck", {"title": "T", "slides": [{"type": "title", "title": "A"}]})]),
        Completion("", [ToolCall("c2", "finish", {"summary": "ok"})]),
    ]
    r = await client.post("/api/decks", json={"brief": "тёмная", "upload_ids": [], "template": "dark"})
    assert r.status_code == 202
    await wait_job(client, r.json()["job_id"])
    d = (await client.get(f"/api/decks/{r.json()['deck_id']}")).json()
    assert d["deck"]["template"] == "dark" and d["template"]["name"] == "Тёмный" and d["has_build"]
    assert (await client.post("/api/decks", json={"brief": "x", "upload_ids": [], "template": "nope"})).status_code == 422

    h = (await client.get("/api/history")).json()
    assert h[0]["text"] == "тёмная" and h[0]["deck_id"] == r.json()["deck_id"] and h[0]["deck_title"] == "T"


def _render_boom(monkeypatch):
    from typst_deck.api.common import ApiError
    from typst_deck.service import DeckService

    def boom(self, deck):
        raise ApiError(500, "render", "шаблон не собрался")

    monkeypatch.setattr(DeckService, "render", boom)


@requires_typst
async def test_failed_render_in_generate_reverts_deck(client, fake_llm, monkeypatch):
    """Провал рендера не оставляет сохранённой колоду, которая не собирается."""
    fake_llm._queue += [
        Completion("", [ToolCall("c1", "set_deck", {"title": "T", "slides": [{"type": "title", "title": "Первый"}]})]),
        Completion("готово", []),
    ]
    r = await client.post("/api/decks", json={"brief": "x", "upload_ids": []})
    deck_id = r.json()["deck_id"]
    assert (await wait_job(client, r.json()["job_id"]))["status"] == "done"
    good = (await client.get(f"/api/decks/{deck_id}")).json()["deck"]

    _render_boom(monkeypatch)
    fake_llm._queue += [
        Completion("", [ToolCall("c2", "set_deck", {"title": "T2", "slides": [{"type": "title", "title": "Второй"}]})]),
        Completion("готово", []),
    ]
    r2 = await client.post(f"/api/decks/{deck_id}/messages", json={"text": "поменяй"})
    assert (await wait_job(client, r2.json()["job_id"]))["status"] == "failed"
    after = (await client.get(f"/api/decks/{deck_id}")).json()["deck"]
    assert [s["title"] for s in after["slides"]] == [s["title"] for s in good["slides"]]


@requires_typst
async def test_failed_template_change_keeps_working_template(client, fake_llm, monkeypatch):
    """Смена шаблона с провалом рендера не должна запирать колоду на несобираемом оформлении."""
    fake_llm._queue += [
        Completion("", [ToolCall("c1", "set_deck", {"title": "T", "slides": [{"type": "title", "title": "A"}]})]),
        Completion("готово", []),
    ]
    r = await client.post("/api/decks", json={"brief": "x", "upload_ids": []})
    deck_id = r.json()["deck_id"]
    assert (await wait_job(client, r.json()["job_id"]))["status"] == "done"

    _render_boom(monkeypatch)
    assert (await client.post(f"/api/decks/{deck_id}/template", json={"template": "dark"})).status_code == 500
    assert (await client.get(f"/api/decks/{deck_id}")).json()["deck"]["template"] == "default"


@requires_typst
async def test_generate_records_tokens_in_chat(client, fake_llm):
    fake_llm._queue += [
        Completion("", [ToolCall("c1", "set_deck", {"title": "T", "slides": [{"type": "title", "title": "A"}]})],
                   tokens_in=1000, tokens_out=200, tokens_cached=800),
        Completion("", [ToolCall("c2", "finish", {"summary": "Готово"})],
                   tokens_in=1500, tokens_out=90, tokens_cached=1400),
    ]
    r = await client.post("/api/decks", json={"brief": "x", "upload_ids": []})
    assert (await wait_job(client, r.json()["job_id"]))["status"] == "done"
    reply = (await client.get(f"/api/decks/{r.json()['deck_id']}/chat")).json()[-1]
    assert reply["role"] == "assistant"
    assert reply["tokens_in"] == 2500 and reply["tokens_out"] == 290
    assert reply["tokens_cached"] == 2200 and reply["iterations"] == 2
    assert reply["llm_ms"] >= 0


@requires_typst
async def test_generate_logs_stages_with_deck_and_job(client, fake_llm, caplog):
    caplog.set_level("INFO", logger="typst_deck.service")
    fake_llm._queue += [
        Completion("", [ToolCall("c1", "set_deck", {"title": "T", "slides": [{"type": "title", "title": "A"}]})]),
        Completion("", [ToolCall("c2", "finish", {"summary": "Готово"})]),
    ]
    r = await client.post("/api/decks", json={"brief": "x", "upload_ids": []})
    deck_id, job_id = r.json()["deck_id"], r.json()["job_id"]
    assert (await wait_job(client, job_id))["status"] == "done"
    events = {rec.msg: rec.fields for rec in caplog.records if hasattr(rec, "fields")}
    assert {"job.start", "stage.reading", "job.done"} <= set(events)
    assert events["job.start"]["deck"] == deck_id and events["job.start"]["job"] == job_id
    assert events["stage.reading"]["files"] == 0
    assert events["job.done"]["v"] == 2 and events["job.done"]["slides"] == 1
    assert events["job.done"]["tok_in"] == 0 and events["job.done"]["ms"] >= 0


@requires_typst
async def test_generate_gives_the_model_template_limits(client, fake_llm, monkeypatch):
    """Сервис обязан передать агенту объявление шаблона колоды, а не пустое."""
    from typst_deck.model import Deck

    seen = {}
    import typst_deck.service as service_module
    real = service_module.run_agent

    async def spy(*a, **kw):
        seen.update(kw)
        return await real(*a, **kw)

    monkeypatch.setattr(service_module, "run_agent", spy)
    store = client.app.state.store
    d = store.create(Deck(title="T", template="dark", slides=[{"type": "final"}]))
    fake_llm._queue += [Completion("", [ToolCall("c1", "finish", {"summary": "готово"})])]
    r = await client.post(f"/api/decks/{d.id}/messages", json={"text": "привет"})
    await wait_job(client, r.json()["job_id"])
    # пределы всех шаблонов, а не только текущего: модель может сменить шаблон посреди прогона
    assert seen["limits"]["dark"] == ({}, "Тёмный")
    assert seen["limits"]["default"] == ({}, "Корпоративный")


@requires_typst
async def test_failed_generation_logs_job_failed(client, fake_llm, caplog):
    caplog.set_level("INFO", logger="typst_deck.service")
    r = await client.post("/api/decks", json={"brief": "x", "upload_ids": []})
    assert (await wait_job(client, r.json()["job_id"]))["status"] == "failed"
    failed = [rec for rec in caplog.records if rec.msg == "job.failed"]
    assert len(failed) == 1
    assert failed[0].fields["deck"] == r.json()["deck_id"]
    assert "LLMError" in failed[0].fields["err"]


@requires_typst
async def test_manual_edit_logs_op(client, fake_llm, caplog):
    fake_llm._queue += [
        Completion("", [ToolCall("c1", "set_deck", {"title": "T", "slides": [{"type": "title", "title": "A"}]})]),
        Completion("", [ToolCall("c2", "finish", {"summary": "Готово"})]),
    ]
    r = await client.post("/api/decks", json={"brief": "x", "upload_ids": []})
    deck_id = r.json()["deck_id"]
    await wait_job(client, r.json()["job_id"])
    caplog.set_level("INFO", logger="typst_deck.service")
    assert (await client.patch(f"/api/decks/{deck_id}", json={"title": "Новое"})).status_code == 200
    done = [x for x in caplog.records if x.msg == "manual.done"]
    assert len(done) == 1
    assert done[0].fields["deck"] == deck_id and done[0].fields["v"] == 3 and done[0].fields["ms"] >= 0


async def test_health_reports_zero_tokens_before_any_call(client):
    h = (await client.get("/api/health")).json()
    assert h["tokens"] == {"in": 0, "out": 0, "cached": 0, "requests": 0}
    assert h["since"]


@requires_typst
async def test_health_accumulates_tokens_across_generations(client, fake_llm):
    fake_llm._queue += [
        Completion("", [ToolCall("c1", "set_deck", {"title": "T", "slides": [{"type": "title", "title": "A"}]})],
                   tokens_in=1000, tokens_out=200, tokens_cached=800),
        Completion("", [ToolCall("c2", "finish", {"summary": "Готово"})],
                   tokens_in=1500, tokens_out=90, tokens_cached=1400),
    ]
    r = await client.post("/api/decks", json={"brief": "x", "upload_ids": []})
    await wait_job(client, r.json()["job_id"])
    h = (await client.get("/api/health")).json()
    assert h["tokens"] == {"in": 2500, "out": 290, "cached": 2200, "requests": 2}


async def test_request_timing_is_logged(client, caplog):
    caplog.set_level("INFO", logger="typst_deck.app")
    await client.get("/api/decks")
    http = [r for r in caplog.records if r.msg == "http"]
    assert len(http) == 1
    assert http[0].fields["method"] == "GET"
    assert http[0].fields["path"] == "/api/decks"
    assert http[0].fields["status"] == 200
    assert http[0].fields["ms"] >= 0


async def test_health_requests_are_not_logged(client, caplog):
    caplog.set_level("INFO", logger="typst_deck.app")
    await client.get("/api/health")
    assert [r for r in caplog.records if r.msg == "http"] == []


async def test_create_app_installs_log_handler(settings):
    import logging as _logging

    from typst_deck.app import create_app
    from typst_deck.logs import HANDLER_MARK

    for h in list(_logging.getLogger().handlers):
        if getattr(h, HANDLER_MARK, False):
            _logging.getLogger().removeHandler(h)
    create_app(settings, llm=None)
    assert any(getattr(h, HANDLER_MARK, False) for h in _logging.getLogger().handlers)


@requires_typst
async def test_render_does_not_block_event_loop(client, fake_llm, monkeypatch):
    """Пока идёт компиляция typst, сервер обязан отвечать на другие запросы."""
    import time

    from typst_deck import service as service_mod

    fake_llm._queue += [
        Completion("", [ToolCall("c1", "set_deck", {"title": "T", "slides": [{"type": "title", "title": "A"}]})]),
        Completion("", [ToolCall("c2", "finish", {"summary": "Готово"})]),
    ]
    r = await client.post("/api/decks", json={"brief": "x", "upload_ids": []})
    deck_id = r.json()["deck_id"]
    await wait_job(client, r.json()["job_id"])

    real_compile = service_mod.compile_deck

    def slow_compile(*a, **kw):
        time.sleep(0.4)  # блокирующий вызов, как настоящий subprocess.run
        return real_compile(*a, **kw)

    monkeypatch.setattr(service_mod, "compile_deck", slow_compile)

    # мерим от старта правки: если цикл занят компиляцией, ответить раньше неё нельзя
    started = time.perf_counter()
    edit = asyncio.create_task(client.patch(f"/api/decks/{deck_id}", json={"title": "Новое"}))
    await asyncio.sleep(0.05)  # даём правке дойти до компиляции
    probe = await client.get("/api/decks")
    probe_s = time.perf_counter() - started

    assert probe.status_code == 200
    assert probe_s < 0.3, f"параллельный запрос ждал рендера {probe_s:.2f} с — цикл заблокирован"
    assert (await edit).status_code == 200


@requires_typst
async def test_template_preview_does_not_block_event_loop(client, monkeypatch):
    """Первое превью шаблона компилирует образец — тоже в потоке, не в цикле."""
    import time

    from typst_deck import service as service_mod

    real_compile = service_mod.compile_deck

    def slow_compile(*a, **kw):
        time.sleep(0.4)
        return real_compile(*a, **kw)

    monkeypatch.setattr(service_mod, "compile_deck", slow_compile)

    started = time.perf_counter()
    preview = asyncio.create_task(client.get("/api/templates/dark/preview/1.svg"))
    await asyncio.sleep(0.05)
    probe = await client.get("/api/decks")
    probe_s = time.perf_counter() - started

    assert probe.status_code == 200
    assert probe_s < 0.3, f"параллельный запрос ждал превью {probe_s:.2f} с — цикл заблокирован"
    assert (await preview).status_code == 200


@requires_typst
async def test_concurrent_first_previews_of_same_template(client):
    """Раньше компиляции сериализовал заблокированный цикл. Теперь они в потоках,
    и два первых превью одного шаблона не должны драться за build.tmp."""
    rs = await asyncio.gather(*(client.get("/api/templates/dark/preview/1.svg") for _ in range(4)))
    assert [r.status_code for r in rs] == [200, 200, 200, 200]
    assert all(r.content.startswith(b"<svg") or b"<svg" in r.content[:200] for r in rs)


@requires_typst
async def test_deck_response_carries_page_size(client, fake_llm):
    """Размер страницы приходит вместе с колодой: фронту незачем качать slide-1.svg ради viewBox."""
    fake_llm._queue += [
        Completion("", [ToolCall("c1", "set_deck", {"title": "T", "slides": [{"type": "title", "title": "A"}]})]),
        Completion("", [ToolCall("c2", "finish", {"summary": "Готово"})]),
    ]
    r = await client.post("/api/decks", json={"brief": "x", "upload_ids": []})
    await wait_job(client, r.json()["job_id"])
    d = (await client.get(f"/api/decks/{r.json()['deck_id']}")).json()
    assert d["page"] == {"w": 453.543307087, "h": 255.118110236}


async def test_deck_response_page_size_without_build(client, fake_llm):
    from typst_deck.render import DEFAULT_PAGE

    fake_llm._queue += [Completion("нет правок", [])]
    r = await client.post("/api/decks", json={"brief": "x", "upload_ids": []})
    await wait_job(client, r.json()["job_id"])
    d = (await client.get(f"/api/decks/{r.json()['deck_id']}")).json()
    assert d["has_build"] is False and d["page"] == DEFAULT_PAGE


@requires_typst
async def test_slide_svg_is_cacheable_by_version(client, fake_llm):
    """URL слайда версионирован ?v=, поэтому содержимое под ним неизменно:
    no-store заставлял браузер качать все слайды заново после каждой правки."""
    fake_llm._queue += [
        Completion("", [ToolCall("c1", "set_deck", {"title": "T", "slides": [{"type": "title", "title": "A"}]})]),
        Completion("", [ToolCall("c2", "finish", {"summary": "Готово"})]),
    ]
    r = await client.post("/api/decks", json={"brief": "x", "upload_ids": []})
    deck_id = r.json()["deck_id"]
    await wait_job(client, r.json()["job_id"])
    v = (await client.get(f"/api/decks/{deck_id}")).json()["version"]

    svg = await client.get(f"/api/decks/{deck_id}/slides/1.svg?v={v}")
    cc = svg.headers["cache-control"]
    assert "no-store" not in cc
    assert "immutable" in cc and "max-age=" in cc
    assert "private" in cc, "колоды за Basic Auth — общим прокси их кэшировать нельзя"
    assert svg.headers["etag"], "ETag нужен, чтобы работал условный запрос"


@requires_typst
async def test_slide_svg_answers_304_on_revalidation(client, fake_llm):
    fake_llm._queue += [
        Completion("", [ToolCall("c1", "set_deck", {"title": "T", "slides": [{"type": "title", "title": "A"}]})]),
        Completion("", [ToolCall("c2", "finish", {"summary": "Готово"})]),
    ]
    r = await client.post("/api/decks", json={"brief": "x", "upload_ids": []})
    deck_id = r.json()["deck_id"]
    await wait_job(client, r.json()["job_id"])

    first = await client.get(f"/api/decks/{deck_id}/slides/1.svg?v=2")
    again = await client.get(f"/api/decks/{deck_id}/slides/1.svg?v=2",
                             headers={"If-None-Match": first.headers["etag"]})
    assert again.status_code == 304
    assert again.content == b""


@requires_typst
async def test_deck_response_carries_per_slide_versions(client, fake_llm):
    """Версия на слайд: правка одного слайда не должна менять URL остальных."""
    fake_llm._queue += [
        Completion("", [ToolCall("c1", "set_deck", {"title": "T", "slides": [
            {"type": "title", "title": "Первый"},
            {"type": "section", "title": "Второй"},
            {"type": "final"}]})]),
        Completion("", [ToolCall("c2", "finish", {"summary": "Готово"})]),
    ]
    r = await client.post("/api/decks", json={"brief": "x", "upload_ids": []})
    deck_id = r.json()["deck_id"]
    await wait_job(client, r.json()["job_id"])

    d = (await client.get(f"/api/decks/{deck_id}")).json()
    before = d["pages"]
    assert len(before) == 3 and len(set(before)) == 3

    sid = d["deck"]["slides"][1]["id"]
    after = (await client.patch(f"/api/decks/{deck_id}/slides/{sid}",
                                json={"patch": {"title": "Изменённый"}})).json()["pages"]
    assert after[1] != before[1]
    assert [after[0], after[2]] == [before[0], before[2]]


async def test_deck_response_pages_empty_without_build(client, fake_llm):
    fake_llm._queue += [Completion("нет правок", [])]
    r = await client.post("/api/decks", json={"brief": "x", "upload_ids": []})
    await wait_job(client, r.json()["job_id"])
    assert (await client.get(f"/api/decks/{r.json()['deck_id']}")).json()["pages"] == []


@requires_typst
async def test_pdf_is_built_on_demand_and_reused(client, fake_llm):
    fake_llm._queue += [
        Completion("", [ToolCall("c1", "set_deck", {"title": "T", "slides": [{"type": "title", "title": "A"}]})]),
        Completion("", [ToolCall("c2", "finish", {"summary": "Готово"})]),
    ]
    r = await client.post("/api/decks", json={"brief": "x", "upload_ids": []})
    deck_id = r.json()["deck_id"]
    await wait_job(client, r.json()["job_id"])

    build = client.app.state.service.store.deck_dir(deck_id) / "build"
    assert not (build / "deck.pdf").exists(), "после сборки колоды PDF ещё не нужен"
    assert (await client.get(f"/api/decks/{deck_id}")).json()["has_build"] is True

    first = await client.get(f"/api/decks/{deck_id}/deck.pdf")
    assert first.status_code == 200 and first.content[:5] == b"%PDF-"
    stamp = (build / "deck.pdf").stat().st_mtime_ns

    second = await client.get(f"/api/decks/{deck_id}/deck.pdf")
    assert second.status_code == 200
    assert (build / "deck.pdf").stat().st_mtime_ns == stamp, "готовый PDF не должен пересобираться"


@requires_typst
async def test_pdf_rebuilt_after_edit(client, fake_llm):
    fake_llm._queue += [
        Completion("", [ToolCall("c1", "set_deck", {"title": "T", "slides": [{"type": "title", "title": "A"}]})]),
        Completion("", [ToolCall("c2", "finish", {"summary": "Готово"})]),
    ]
    r = await client.post("/api/decks", json={"brief": "x", "upload_ids": []})
    deck_id = r.json()["deck_id"]
    await wait_job(client, r.json()["job_id"])
    await client.get(f"/api/decks/{deck_id}/deck.pdf")

    await client.patch(f"/api/decks/{deck_id}", json={"title": "Другое имя"})
    build = client.app.state.service.store.deck_dir(deck_id) / "build"
    assert not (build / "deck.pdf").exists(), "правка обязана обнулить готовый PDF"
    assert (await client.get(f"/api/decks/{deck_id}/deck.pdf")).content[:5] == b"%PDF-"


@requires_typst
async def test_pdf_build_is_logged(client, fake_llm, caplog):
    fake_llm._queue += [
        Completion("", [ToolCall("c1", "set_deck", {"title": "T", "slides": [{"type": "title", "title": "A"}]})]),
        Completion("", [ToolCall("c2", "finish", {"summary": "Готово"})]),
    ]
    r = await client.post("/api/decks", json={"brief": "x", "upload_ids": []})
    deck_id = r.json()["deck_id"]
    await wait_job(client, r.json()["job_id"])
    caplog.set_level("INFO", logger="typst_deck.render")
    await client.get(f"/api/decks/{deck_id}/deck.pdf")
    done = [x for x in caplog.records if x.msg == "pdf.done"]
    assert len(done) == 1 and done[0].fields["ms"] >= 0 and done[0].fields["kb"] > 0


@requires_typst
async def test_failed_render_is_reported_to_model_on_retry(client, fake_llm, monkeypatch):
    """Без этого модель повторяет ту же правку: она не знает, что прошлая не собралась."""
    from typst_deck import service as service_mod
    from typst_deck.render import CompileError

    fake_llm._queue += [
        Completion("", [ToolCall("c1", "set_deck", {"title": "T", "slides": [{"type": "title", "title": "A"}]})]),
        Completion("", [ToolCall("c2", "finish", {"summary": "Готово"})]),
    ]
    r = await client.post("/api/decks", json={"brief": "первый заход", "upload_ids": []})
    deck_id = r.json()["deck_id"]
    await wait_job(client, r.json()["job_id"])

    def boom(*a, **kw):
        raise CompileError("слайд 1 «A» не помещается на страницу: содержимого больше")

    monkeypatch.setattr(service_mod, "compile_deck", boom)
    fake_llm._queue += [
        Completion("", [ToolCall("c3", "set_deck", {"title": "T", "slides": [
            {"type": "title", "title": "Длинный заголовок"}]})]),
        Completion("", [ToolCall("c4", "finish", {"summary": "ок"})]),
    ]
    j = await client.post(f"/api/decks/{deck_id}/messages", json={"text": "добавь отступы"})
    assert (await wait_job(client, j.json()["job_id"]))["status"] == "failed"

    monkeypatch.undo()
    fake_llm.calls.clear()
    fake_llm._queue += [Completion("теперь короче", [])]
    j = await client.post(f"/api/decks/{deck_id}/messages", json={"text": "давай ещё раз"})
    await wait_job(client, j.json()["job_id"])

    sent = fake_llm.calls[0]
    assert any("не помещается на страницу" in (m.get("content") or "") for m in sent), \
        [m.get("content", "")[:60] for m in sent]


@requires_typst
async def test_manual_edit_over_limit_is_422_with_the_same_text(tight_client):
    from typst_deck.model import Deck

    store = tight_client.app.state.store
    d = store.create(Deck(title="T", template="tight", slides=[{"type": "title", "title": "Итоги"}]))
    sid = d.slides[0].id
    r = await tight_client.patch(f"/api/decks/{d.id}/slides/{sid}", json={"patch": {"title": "Э" * 80}})
    assert r.status_code == 422
    msg = r.json()["error"]["message"]
    assert "«Заголовок»" in msg and "80 знаков при пределе 60" in msg and "Тесный" in msg
    assert (await tight_client.get(f"/api/decks/{d.id}")).json()["deck"]["slides"][0]["title"] == "Итоги"


@requires_typst
async def test_manual_edit_that_reduces_an_old_violation_passes(tight_client):
    """Колода, созданная до появления пределов, обязана оставаться редактируемой."""
    from typst_deck.model import Deck

    store = tight_client.app.state.store
    d = store.create(Deck(title="T", template="tight", slides=[{"type": "title", "title": "Э" * 110}]))
    sid = d.slides[0].id
    r = await tight_client.patch(f"/api/decks/{d.id}/slides/{sid}", json={"patch": {"title": "Э" * 100}})
    assert r.status_code == 200
    r = await tight_client.patch(f"/api/decks/{d.id}/slides/{sid}", json={"patch": {"title": "Э" * 105}})
    assert r.status_code == 422


@requires_typst
async def test_generate_checks_limits_of_the_template_the_model_switched_to(tight_client, fake_llm):
    """set_template("tight") и правка в том же прогоне: правка проверяется пределами
    нового шаблона, иначе модель пишет сверх ёмкости в один запрос."""
    from typst_deck.model import Deck

    store = tight_client.app.state.store
    d = store.create(Deck(title="T", template="default", slides=[{"type": "title", "title": "Итоги"}]))
    sid = d.slides[0].id
    fake_llm._queue += [
        Completion("", [ToolCall("c1", "set_template", {"template": "tight"})]),
        Completion("", [ToolCall("c2", "update_slide", {"slide_id": sid, "patch": {"title": "Э" * 80}})]),
        Completion("", [ToolCall("c3", "finish", {"summary": "готово"})]),
    ]
    r = await tight_client.post(f"/api/decks/{d.id}/messages", json={"text": "в тесный и длиннее"})
    assert (await wait_job(tight_client, r.json()["job_id"]))["status"] == "done"
    error = fake_llm.calls[2][-1]["content"]
    assert error.startswith("ERROR") and "Тесный" in error, error
    deck = (await tight_client.get(f"/api/decks/{d.id}")).json()["deck"]
    assert deck["template"] == "tight" and deck["slides"][0]["title"] == "Итоги"


@requires_typst
async def test_revert_is_not_blocked_by_limits(tight_client):
    """Откат возвращает состояние, которое уже собиралось; запирать его нельзя."""
    from typst_deck.model import Deck

    store = tight_client.app.state.store
    d = store.create(Deck(title="T", template="tight", slides=[{"type": "title", "title": "Э" * 110}]))
    sid = d.slides[0].id
    assert (await tight_client.patch(f"/api/decks/{d.id}/slides/{sid}",
                                     json={"patch": {"title": "Коротко"}})).status_code == 200
    v = (await tight_client.get(f"/api/decks/{d.id}/versions")).json()[0]["version"]
    r = await tight_client.post(f"/api/decks/{d.id}/revert", json={"version": v})
    assert r.status_code == 200


@requires_typst
async def test_rename_and_move_are_not_affected_by_limits(tight_client):
    from typst_deck.model import Deck

    store = tight_client.app.state.store
    d = store.create(Deck(title="T", template="tight",
                          slides=[{"type": "title", "title": "Э" * 110}, {"type": "final"}]))
    assert (await tight_client.patch(f"/api/decks/{d.id}", json={"title": "Новое"})).status_code == 200
    sid = d.slides[1].id
    assert (await tight_client.post(f"/api/decks/{d.id}/slides/{sid}/move",
                                    json={"to_index": 0})).status_code == 200


@requires_typst
async def test_moving_to_a_tighter_template_warns_but_is_allowed(tight_client):
    """Блокировать нельзя: колода застрянет — ни остаться, ни переехать."""
    from typst_deck.model import Deck

    store = tight_client.app.state.store
    d = store.create(Deck(title="T", template="default", slides=[{"type": "title", "title": "Э" * 100}]))
    r = await tight_client.post(f"/api/decks/{d.id}/template", json={"template": "tight"})
    assert r.status_code == 200 and r.json()["deck"]["template"] == "tight"
    assert any("«Заголовок»" in w and "Тесный" in w for w in r.json()["warnings"])
    # предупреждение живёт, пока живёт причина, и снимается вместе с ней
    sid = d.slides[0].id
    r = await tight_client.patch(f"/api/decks/{d.id}/slides/{sid}", json={"patch": {"title": "Коротко"}})
    assert r.status_code == 200 and r.json()["warnings"] == []


async def test_limit_warnings_on_an_old_deck_are_capped(tight_client):
    """Колода, созданная до пределов, не должна давать стену текста на каждом чтении."""
    from typst_deck.model import Deck

    store = tight_client.app.state.store
    d = store.create(Deck(title="T", template="tight", slides=[{"type": "title", "title": "Э" * 80}] * 7))
    warnings = (await tight_client.get(f"/api/decks/{d.id}")).json()["warnings"]
    assert len(warnings) == 6 and warnings[-1] == "и ещё 2", warnings
    assert warnings[0].startswith("слайд 1, поле «Заголовок»")


@requires_typst
async def test_deck_without_declared_slots_has_no_such_warnings(client):
    from typst_deck.model import Deck

    store = client.app.state.store
    d = store.create(Deck(title="T", slides=[{"type": "title", "title": "Э" * 110}]))
    r = await client.post(f"/api/decks/{d.id}/template", json={"template": "dark"})
    assert r.status_code == 200 and r.json()["warnings"] == []


async def test_pdf_is_built_with_the_template_fonts(settings, fake_llm, tmp_path, monkeypatch):
    """PDF собирается с каталогом шрифтов шаблона колоды, а не с жёстким build/template/fonts."""
    import json
    import shutil

    import typst_deck.service as service_module
    from typst_deck.app import create_app
    from typst_deck.model import Deck

    root = tmp_path / "corp"
    t = root / "withfonts"
    shutil.copytree(settings.template_dir, t)
    (t / "typst-deck" / "fonts").mkdir(parents=True)
    (t / "template.json").write_text(json.dumps({"name": "Со шрифтами", "fonts": "typst-deck/fonts"}))
    app = create_app(settings.model_copy(update={"templates_dir": str(root)}), llm=fake_llm)
    store, svc_ = app.state.store, app.state.service
    d = store.create(Deck(title="T", template="withfonts", slides=[{"type": "final"}]))
    build = store.deck_dir(d.id) / "build"
    build.mkdir()
    (build / "deck.typ").write_text("")
    (build / "template").symlink_to(t.resolve())  # собрано под действующий шаблон: пересобирать нечего
    seen = {}

    def fake_pdf(build_dir, typst_bin, fonts_dir=None):
        seen["fonts"] = fonts_dir
        (build_dir / "deck.pdf").write_bytes(b"%PDF")
        return build_dir / "deck.pdf"

    monkeypatch.setattr(service_module, "compile_pdf", fake_pdf)
    await svc_.pdf(d.id)
    assert seen["fonts"] == t / "typst-deck" / "fonts"


async def test_pdf_of_a_deck_whose_template_vanished_is_refused(client):
    """Шаблон колоды пропал — PDF не собирается молча с build/template/fonts, а называет причину."""
    from typst_deck.model import Deck

    store = client.app.state.store
    d = store.create(Deck(title="T", template="gone", slides=[{"type": "final"}]))
    build = store.deck_dir(d.id) / "build"
    build.mkdir()
    (build / "deck.typ").write_text("")
    r = await client.get(f"/api/decks/{d.id}/deck.pdf")
    assert r.status_code == 422 and "не найден" in r.json()["error"]["message"]


@requires_typst
@pytest.mark.parametrize("old_left", [False, True], ids=["old-gone", "old-left"])
async def test_pdf_rebuilds_a_build_left_on_the_old_template_path(settings, fake_llm, tmp_path, old_left):
    """Колоду перенесли на новый id шаблона, а её build/template смотрит на прежний
    каталог: в новом образе его нет (old-gone) или он уже не действующий (old-left).
    PDF обязан пересобрать build/ под текущий шаблон, а не отдать 500 «file not found»,
    и при этом не порождать правку: ни новой версии, ни записи в чате."""
    import shutil

    from tests.conftest import make_template
    from typst_deck.app import create_app
    from typst_deck.migrate_template import migrate
    from typst_deck.model import Deck

    old = make_template(tmp_path / "old-root", "old")
    new = make_template(tmp_path / "new-root", "new")
    before = create_app(settings.model_copy(update={"templates_dir": str(old.parent)}), llm=fake_llm)
    d = before.state.store.create(Deck(title="T", template="old", slides=[{"type": "final"}]))
    await before.state.service.render(d)
    build = before.state.store.deck_dir(d.id) / "build"
    assert (build / "template").resolve() == old.resolve()
    assert not (build / "deck.pdf").exists()

    if not old_left:
        shutil.rmtree(old)
    assert migrate(before.state.store, "old", "new") == [d.id]
    saved = before.state.store.load(d.id)
    app = create_app(settings.model_copy(update={"templates_dir": str(new.parent)}), llm=fake_llm)
    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://t") as c:
        r = await c.get(f"/api/decks/{d.id}/deck.pdf")
    assert r.status_code == 200, r.text
    assert r.content[:5] == b"%PDF-"
    assert (build / "template").resolve() == new.resolve()
    after = app.state.store.load(d.id)
    assert (after.version, after.updated_at) == (saved.version, saved.updated_at), "скачивание PDF — не правка"
    assert app.state.store.read_chat(d.id) == []


@requires_typst
async def test_old_template_id_is_stored_as_the_new_one(settings, fake_llm):
    """Бот мог зашить старый id: создание и переезд принимают его и пишут новый."""
    import httpx

    from typst_deck.app import create_app
    from typst_deck.model import Deck

    app = create_app(settings.model_copy(update={"template_aliases": "old-dark:dark"}), llm=fake_llm)
    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://t") as c:
        fake_llm._queue += [Completion("", [ToolCall("c1", "finish", {"summary": "ok"})])]
        r = await c.post("/api/decks", json={"brief": "Итоги", "template": "old-dark"})
        assert r.status_code == 202
        await wait_job(c, r.json()["job_id"])
        assert (await c.get(f"/api/decks/{r.json()['deck_id']}")).json()["deck"]["template"] == "dark"
        assert "old-dark" not in [t["id"] for t in (await c.get("/api/templates")).json()]
        d = app.state.store.create(Deck(title="T", slides=[{"type": "final"}]))
        r = await c.post(f"/api/decks/{d.id}/template", json={"template": "old-dark"})
        assert r.status_code == 200 and r.json()["deck"]["template"] == "dark"


@requires_typst
async def test_template_preview_by_old_id_uses_the_cache_of_the_new(settings, fake_llm):
    """Старый id из алиаса не заводит свой кэш превью рядом с действующим."""
    from typst_deck.app import create_app

    app = create_app(settings.model_copy(update={"template_aliases": "old-dark:dark"}), llm=fake_llm)
    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://t") as c:
        assert (await c.get("/api/templates/old-dark/preview/1.svg")).status_code == 200
    assert [p.name for p in (settings.data_dir / "template-previews").iterdir()] == ["dark"]


@requires_typst
async def test_broken_template_is_hidden_refused_and_reported(settings, fake_llm, tmp_path):
    import httpx

    from tests.conftest import TWO_PAGES, make_template
    from typst_deck.app import create_app

    root = tmp_path / "corp"
    make_template(root, "broken", deck_edit=TWO_PAGES, meta={"name": "Сломанный"})
    app = create_app(settings.model_copy(update={"templates_dir": str(root)}), llm=fake_llm)
    app.state.service.check_templates()
    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://t") as c:
        assert "broken" not in [t["id"] for t in (await c.get("/api/templates")).json()]
        r = await c.post("/api/decks", json={"brief": "Итоги", "template": "broken"})
        assert r.status_code == 422 and "сломан" in r.json()["error"]["message"]
        h = (await c.get("/api/health")).json()
        assert h["ok"] is True and h["templates"]["checked"] is True
        assert "broken" in h["templates"]["broken"] and "default" in h["templates"]["ok"]
        assert h["templates"]["error"] is None  # проверка завершилась, а не упала


async def test_check_templates_survives_an_exception_and_remembers_it(settings, fake_llm, monkeypatch):
    """check_templates идёт в фоновом потоке (app.py): необработанное исключение
    там тихо роняет поток — checked навсегда остаётся False, а сломанный шаблон
    не скрывается и не отвергается. Поэтому падение самой проверки не должно
    убивать вызов, а причина должна быть видна в /api/health."""
    from typst_deck.app import create_app

    app = create_app(settings, llm=fake_llm)
    svc = app.state.service

    def boom(typst_bin, work_dir):
        raise RuntimeError("бабах")

    monkeypatch.setattr(svc.templates, "check_all", boom)
    svc.check_templates()  # не должно бросать исключение наружу
    health = svc.templates_health()
    assert health["checked"] is False  # отличимо от «ещё выполняется» через error
    assert health["error"] is not None and "бабах" in health["error"]


@requires_typst
async def test_app_checks_templates_in_background_on_start(settings, fake_llm):
    from typst_deck.app import create_app

    app = create_app(settings.model_copy(update={"check_templates_on_start": True}), llm=fake_llm)
    for _ in range(200):
        if app.state.service.templates.checked:
            break
        await asyncio.sleep(0.05)
    assert app.state.service.templates.checked


async def test_health_names_templates_with_pptx_off(client):
    client.app.state.service.templates.pptx_off["x"] = "в мастере нет макета «Y»"
    t = (await client.get("/api/health")).json()["templates"]
    assert t["pptx_off"] == {"x": "в мастере нет макета «Y»"}


PPTX_MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"


def put_deck(client, template="withpptx"):
    return client.app.state.store.create(Deck(title="Итоги квартала", template=template, slides=[
        {"type": "title", "title": "Итоги"}, {"type": "final"}]))


async def test_pptx_download_changes_nothing(pptx_client):
    deck = put_deck(pptx_client)
    r = await pptx_client.get(f"/api/decks/{deck.id}/deck.pptx")
    assert r.status_code == 200 and r.headers["content-type"] == PPTX_MIME
    assert r.headers["content-disposition"] == "attachment; filename*=utf-8''" + quote("Итоги квартала.pptx")
    assert len(Presentation(io.BytesIO(r.content)).slides) == 2
    d = (await pptx_client.get(f"/api/decks/{deck.id}")).json()
    assert d["version"] == deck.version, "скачивание PPTX не даёт новой версии"
    assert (await pptx_client.get(f"/api/decks/{deck.id}/chat")).json() == []


async def test_pptx_url_and_flags_follow_the_master(pptx_client):
    with_master, without = put_deck(pptx_client), put_deck(pptx_client, template="default")
    d = (await pptx_client.get(f"/api/decks/{with_master.id}")).json()
    assert d["pptx_url"] == f"http://t/api/decks/{with_master.id}/deck.pptx" and d["template"]["pptx"] is True
    assert (await pptx_client.get(f"/api/decks/{without.id}")).json()["pptx_url"] is None
    flags = {x["id"]: x["pptx"] for x in (await pptx_client.get("/api/decks")).json()}
    assert flags == {with_master.id: True, without.id: False}
    tpl = {t["id"]: t["pptx"] for t in (await pptx_client.get("/api/templates")).json()}
    assert tpl == {"default": False, "withpptx": True}
    r = await pptx_client.get(f"/api/decks/{without.id}/deck.pptx")
    assert r.status_code == 422 and "нет PowerPoint-мастера" in r.json()["error"]["message"]


async def test_pptx_off_and_broken_mapping_are_422_with_reason(pptx_client):
    svc = pptx_client.app.state.service
    deck = put_deck(pptx_client)
    svc.templates.pptx_off["withpptx"] = "в мастере нет макета «X»"
    r = await pptx_client.get(f"/api/decks/{deck.id}/deck.pptx")
    assert r.status_code == 422 and "выключен: в мастере нет макета «X»" in r.json()["error"]["message"]
    svc.templates.pptx_off.clear()
    mp = svc.templates.get("withpptx").path / "pptx.json"
    m = json.loads(mp.read_text())
    m["slides"]["final"]["layout"] = "Пропал"
    mp.write_text(json.dumps(m, ensure_ascii=False))
    r = await pptx_client.get(f"/api/decks/{deck.id}/deck.pptx")
    assert r.status_code == 422 and "нет макета «Пропал»" in r.json()["error"]["message"]
    assert (await pptx_client.get("/api/decks/nope/deck.pptx")).status_code == 404


async def test_pptx_export_failure_is_500(pptx_client, monkeypatch):
    import typst_deck.service as service

    def boom(*args, **kwargs):
        raise RuntimeError("битый мастер")

    monkeypatch.setattr(service, "export_pptx", boom)
    r = await pptx_client.get(f"/api/decks/{put_deck(pptx_client).id}/deck.pptx")
    assert r.status_code == 500 and r.json()["error"]["code"] == "render"
    assert "битый мастер" in r.json()["error"]["message"]
