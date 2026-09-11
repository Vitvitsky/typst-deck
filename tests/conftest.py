"""Общие фикстуры тестов."""
import base64
import os
import shutil
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE_DIR = ROOT / "template"
HAS_TYPST = shutil.which("typst") is not None
requires_typst = pytest.mark.skipif(not HAS_TYPST, reason="typst не установлен")

NO_CORP = ("нет корпоративных шаблонов: задайте TYPST_DECK_TEST_TEMPLATES "
           "(корни через «:», как TYPST_DECK_TEMPLATES_DIR)")


def corp_templates() -> list:
    """Корпоративные шаблоны для общих тестов — все шаблоны корней из TYPST_DECK_TEST_TEMPLATES."""
    raw = os.environ.get("TYPST_DECK_TEST_TEMPLATES", "")
    if not raw:
        return []
    from typst_deck.templates import TemplateRegistry

    roots = [Path(p).expanduser() for p in raw.split(os.pathsep) if p]
    return [t for t in TemplateRegistry(TEMPLATE_DIR, roots).list() if t.id != "default"]


CORP_PARAMS = ([pytest.param(t, id=t.id) for t in corp_templates()]
               or [pytest.param(None, id="no-corp", marks=pytest.mark.skip(reason=NO_CORP))])

PNG_1PX = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNkYPhfDwAChwGA60e6kgAAAABJRU5ErkJggg=="
)


# Правки эталона template/deck.typ: каждая ломает ровно одно правило контракта.
TWO_PAGES = ("#let slide-final(title: none, contact: none) = slide(fill: accent)[",
             "#let slide-final(title: none, contact: none) = page[] + slide(fill: accent)[")
NO_CLIP_FINAL = ("#let slide-final(title: none, contact: none) = slide(fill: accent)[",
                 "#let slide-final(title: none, contact: none) = page(fill: accent)[")
NO_CONTACT = ("#if contact != none { v(0.4cm); text(size: 12pt, contact) }", "")
FALLBACK_FONT = ('set text(font: "DejaVu Sans",', 'set text(font: ("DejaVu Sans", "Нет Такого Шрифта"),')


def make_template(parent: Path, name: str, deck_edit: tuple[str, str] | None = None,
                  meta: dict | None = None) -> Path:
    """Копия эталона template/ с одной правкой deck.typ и своим template.json.

    Сломанные шаблоны выводятся из эталона, а не лежат файлами: так они не
    расходятся с ним, а правка, которая перестала применяться, валит тест сразу.
    """
    import json

    d = parent / name
    shutil.copytree(TEMPLATE_DIR, d)
    if deck_edit:
        old, new = deck_edit
        src = (d / "deck.typ").read_text()
        assert old in src, f"эталон template/deck.typ изменился: нет «{old}»"
        (d / "deck.typ").write_text(src.replace(old, new, 1))
    if meta is not None:
        (d / "template.json").write_text(json.dumps(meta, ensure_ascii=False))
    return d


# Сопоставление синтетического мастера: макеты стандартного мастера python-pptx.
# «Title and Content» у image — нарочно: его idx 1 не сопоставлен, на нём
# проверяется удаление лишнего заполнителя.
PPTX_MAPPING = {
    "master": "master.pptx",
    "text": {"size": 24, "bullet": "FFDE00", "caption": "666666"},
    "chart": {"palette": ["FFDE00", "1E1E1E", "00FFE0"], "negative": "FF2E00"},
    "slides": {
        "title": {"layout": "Title Slide", "ph": {"title": 0, "subtitle": 1, "author": 21, "department": 21}},
        "section": {"layout": "Section Header", "ph": {"title": 0, "subtitle": 1}},
        "image": {"layout": "Title and Content", "ph": {"title": 0}, "body": [20, 80, 600, 260],
                  "image_width": {"small": 30, "medium": 42, "large": 55}, "gap": 15},
        "bar-chart": {"layout": "Title Only", "ph": {"title": 0}, "body": [20, 80, 600, 260],
                      "text_width": 38, "gap": 12, "heading": "Выводы"},
        "chart": {"layout": "Title Only", "ph": {"title": 0}, "body": [20, 80, 600, 260],
                  "text_width": 38, "gap": 12},
        "final": {"layout": "Title Slide", "ph": {"title": 0, "contact": 1}},
    },
}


def build_master(path: Path) -> None:
    """Синтетический PowerPoint-мастер: стандартный мастер python-pptx, слайд 677×381 мм
    (нестандартный размер ловит зашитый 4:3), в «Title Slide» — заполнитель idx 21,
    в презентации — один слайд-образец, который экспорт обязан убрать.
    Корпоративные мастера в репозиторий не попадают."""
    import copy

    from pptx import Presentation
    from pptx.oxml.ns import qn
    from pptx.util import Mm

    prs = Presentation()
    prs.slide_width, prs.slide_height = Mm(677), Mm(381)
    layout = next(lay for lay in prs.slide_layouts if lay.name == "Title Slide")
    sub = next(ph for ph in layout.placeholders if ph.placeholder_format.idx == 1)
    el = copy.deepcopy(sub._element)
    cnv = el.find(qn("p:nvSpPr")).find(qn("p:cNvPr"))
    cnv.set("id", "100")
    cnv.set("name", "Author and Date")
    ph = el.find(".//" + qn("p:ph"))
    ph.set("type", "body")
    ph.set("idx", "21")
    layout.shapes._spTree.append(el)
    prs.slides.add_slide(layout)
    prs.save(path)


def make_pptx_template(parent: Path, name: str = "withpptx", mapping_edit=None) -> Path:
    """Копия эталона template/ с синтетическим мастером и сопоставлением pptx.json.

    mapping_edit(mapping) правит хорошее сопоставление перед записью: сломанные
    варианты выводятся из рабочего и не расходятся с ним.
    """
    import copy
    import json

    d = make_template(parent, name, meta={"name": "С мастером", "pptx": "pptx.json"})
    build_master(d / "master.pptx")
    mapping = copy.deepcopy(PPTX_MAPPING)
    if mapping_edit is not None:
        mapping_edit(mapping)
    (d / "pptx.json").write_text(json.dumps(mapping, ensure_ascii=False))
    return d


@pytest.fixture
def deck_dir(tmp_path: Path) -> Path:
    d = tmp_path / "deck"
    (d / "assets").mkdir(parents=True)
    return d


@pytest.fixture
def png_bytes() -> bytes:
    return PNG_1PX


@pytest.fixture
def settings(tmp_path):
    from typst_deck.config import Settings

    return Settings(data_dir=tmp_path / "data", template_dir=TEMPLATE_DIR, templates_dir=str(ROOT / "templates"),
                    public_base_url="http://t", llm_model="deepseek-chat", check_templates_on_start=False)


@pytest.fixture
def fake_llm():
    from typst_deck.llm.client import FakeLLM

    return FakeLLM([])


@pytest.fixture
async def client(settings, fake_llm):
    import httpx

    from typst_deck.app import create_app

    app = create_app(settings, llm=fake_llm)
    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://t") as c:
        c.app = app
        yield c


@pytest.fixture
def tight_settings(tmp_path):
    """Клиент с шаблоном, у которого объявлена ёмкость слотов.

    Шаблон — копия default: важна не вёрстка, а объявление в template.json.
    """
    import json
    import shutil as sh

    from typst_deck.config import Settings

    root = tmp_path / "templates"
    tight = root / "tight"
    tight.mkdir(parents=True)
    sh.copy(TEMPLATE_DIR / "deck.typ", tight / "deck.typ")
    (tight / "template.json").write_text(json.dumps({
        "name": "Тесный",
        "slots": {"title": {"title": {"max": 60}, "subtitle": {"lines": 2}},
                  "image": {"bullets": {"count": 3}}},
    }, ensure_ascii=False))
    return Settings(data_dir=tmp_path / "data", template_dir=TEMPLATE_DIR, templates_dir=str(root),
                    public_base_url="http://t", llm_model="deepseek-chat", check_templates_on_start=False)


@pytest.fixture
async def tight_client(tight_settings, fake_llm):
    import httpx

    from typst_deck.app import create_app

    app = create_app(tight_settings, llm=fake_llm)
    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://t") as c:
        c.app = app
        yield c
