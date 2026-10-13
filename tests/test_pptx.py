"""Экспорт в PPTX: сопоставление шаблона и сборка файла на синтетическом мастере."""
import io
import re

import pytest
from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Mm, Pt

from tests.conftest import make_pptx_template, requires_typst
from typst_deck.contract import SAMPLE_DECK, filled_slide
from typst_deck.model import Deck
from typst_deck.pptx import SLIDE_TYPES, MappingError, export_pptx, load_mapping, open_master

MM = 36000
ZONE = (20, 80, 600, 260)  # body синтетического сопоставления, мм


# --- экспорт

@pytest.fixture
def tdir(tmp_path):
    return make_pptx_template(tmp_path)


def export(tdir, deck, assets=None, warnings=None):
    data = export_pptx(deck, tdir, "pptx.json", assets or tdir / "no-assets", "typst", warnings)
    return Presentation(io.BytesIO(data))


def by_idx(slide) -> dict[int, list[str]]:
    return {sh.placeholder_format.idx: [p.text for p in sh.text_frame.paragraphs] for sh in slide.placeholders}


def pictures(slide):
    return [sh for sh in slide.shapes if sh.shape_type == MSO_SHAPE_TYPE.PICTURE]


def boxes(slide):
    return [sh for sh in slide.shapes if sh.shape_type == MSO_SHAPE_TYPE.TEXT_BOX]


def chart_of(slide):
    return next(sh.chart for sh in slide.shapes if sh.has_chart)


def colors(chart, path: str) -> list[str]:
    return [e.get("val") for e in chart._chartSpace.xpath(path)]


def inside(sh, zone=ZONE) -> bool:
    x, y, w, h = (int(v * MM) for v in zone)
    return (sh.left >= x - 1 and sh.top >= y - 1
            and sh.left + sh.width <= x + w + 1 and sh.top + sh.height <= y + h + 1)


def png(path, w=400, h=200):
    Image.new("RGB", (w, h), (255, 222, 0)).save(path)


def image_deck(size="medium", **kw) -> Deck:
    return Deck(slides=[{"type": "image", "title": "Рост", "bullets": ["Тезис"], "image": "pic.png",
                         "image_size": size, "caption": "Подпись", **kw}])


def test_every_type_lands_on_its_layout_at_the_masters_size(tdir):
    from tests.conftest import PPTX_MAPPING

    prs = export(tdir, SAMPLE_DECK)
    assert [s.slide_layout.name for s in prs.slides] == [
        PPTX_MAPPING["slides"][s.type]["layout"] for s in SAMPLE_DECK.slides]
    assert len(prs.slides) == 15
    assert (prs.slide_width, prs.slide_height) == (Mm(677), Mm(381))
    # слайд-образец мастера не пережил экспорт: первый слайд — титул колоды
    assert by_idx(prs.slides[0])[0] == ["Название презентации"]


def test_fields_fill_placeholders_and_share_one_in_model_order(tdir):
    title = by_idx(export(tdir, SAMPLE_DECK).slides[0])
    assert title[1] == ["Подзаголовок и дата"]
    assert title[21] == ["Имя Фамилия", "Подразделение"]


def test_empty_and_unmapped_placeholders_are_removed(tdir, tmp_path):
    (tmp_path / "a").mkdir()
    png(tmp_path / "a" / "pic.png")
    deck = Deck(slides=[{"type": "title", "title": "Итоги"},
                        {"type": "image", "title": "Рост", "image": "pic.png"}])
    prs = export(tdir, deck, assets=tmp_path / "a")
    # ни подзаголовка, ни подписи, ни idx 1 «Title and Content»; дату и колонтитул python-pptx не переносит
    assert set(by_idx(prs.slides[0])) == {0, 12}
    assert set(by_idx(prs.slides[1])) == {0, 12}


def test_slide_number_is_a_field(tdir):
    for n, slide in enumerate(export(tdir, SAMPLE_DECK).slides, 1):
        nums = [sh for sh in slide.placeholders if sh.placeholder_format.idx == 12]
        assert len(nums) == 1, n  # у «Blank» номер тоже есть
        assert nums[0]._element.xpath(".//a:fld[@type='slidenum']"), n
        assert nums[0].text_frame.text == str(n)


def test_titles_shrink_on_edit(tdir):
    for slide in export(tdir, SAMPLE_DECK).slides:
        if slide.shapes.title is not None:  # у цитаты заголовка нет
            assert slide.shapes.title._element.xpath(".//a:bodyPr/a:normAutofit")


def test_image_keeps_proportions_inside_the_zone(tdir, tmp_path):
    (tmp_path / "assets").mkdir()
    png(tmp_path / "assets" / "pic.png")
    widths = []
    for size in ("small", "medium", "large"):
        slide = export(tdir, image_deck(size), assets=tmp_path / "assets").slides[0]
        [pic] = pictures(slide)
        text, caption = boxes(slide)
        assert inside(pic) and inside(text) and inside(caption)
        assert abs(pic.width / pic.height - 2) < 0.01
        assert text.left + text.width <= pic.left  # тезисы слева от картинки
        assert caption.top >= pic.top + pic.height and caption.text_frame.text == "Подпись"
        widths.append(pic.width)
    assert widths == sorted(widths) and len(set(widths)) == 3


def test_image_takes_exactly_its_share_and_bullets_absorb_the_gap(tdir, tmp_path):
    (tmp_path / "assets").mkdir()
    png(tmp_path / "assets" / "pic.png", 400, 100)
    slide = export(tdir, image_deck("medium"), assets=tmp_path / "assets").slides[0]
    [pic] = pictures(slide)
    text, _caption = boxes(slide)
    zone_w = int(ZONE[2] * MM)
    img_w = int(zone_w * 42 / 100)
    assert abs(pic.width - img_w) <= 1  # картинка — ровно 42 % зоны
    assert text.width == zone_w - img_w - 15 * MM  # зазор вычтен из тезисов


def test_missing_image_leaves_bullets_the_whole_zone(tdir, tmp_path):
    warnings: list[str] = []
    slide = export(tdir, image_deck(), assets=tmp_path, warnings=warnings).slides[0]
    assert pictures(slide) == [] and "pic.png не найдена" in warnings[0]
    [text] = boxes(slide)
    assert (text.left, text.width) == (int(ZONE[0] * MM), int(ZONE[2] * MM))


def test_unreadable_image_leaves_bullets_the_whole_zone(tdir, tmp_path):
    (tmp_path / "assets").mkdir()
    (tmp_path / "assets" / "pic.png").write_bytes(b"not an image")
    warnings: list[str] = []
    slide = export(tdir, image_deck(), assets=tmp_path / "assets", warnings=warnings).slides[0]
    assert pictures(slide) == [] and "pic.png не читается" in warnings[0]
    [text] = boxes(slide)
    assert (text.left, text.width) == (int(ZONE[0] * MM), int(ZONE[2] * MM))


def test_image_cannot_escape_assets(tdir, tmp_path):
    (tmp_path / "assets").mkdir()
    png(tmp_path / "secret.png")
    warnings: list[str] = []
    slide = export(tdir, image_deck(image="../secret.png"), assets=tmp_path / "assets",
                   warnings=warnings).slides[0]
    assert pictures(slide) == [] and "не найдена" in warnings[0]


@requires_typst
def test_svg_goes_in_as_png(tdir, tmp_path):
    (tmp_path / "vec.svg").write_text('<svg xmlns="http://www.w3.org/2000/svg" width="300" height="100">'
                                      '<rect width="300" height="100" fill="#00FFE0"/></svg>')
    slide = export(tdir, image_deck(image="vec.svg"), assets=tmp_path).slides[0]
    [pic] = pictures(slide)
    assert pic.image.content_type == "image/png" and abs(pic.width / pic.height - 3) < 0.05


def bar(**kw) -> dict:
    return {"type": "chart", "kind": "bar", "title": "Итоги", "unit": "млн",
            "labels": ["Центр", "Север", "Юг"], "series": [{"name": "млн", "values": [42, -17, 28]}]} | kw


def test_bar_colors_follow_the_sign(tdir):
    chart = chart_of(export(tdir, Deck(slides=[bar()])).slides[0])
    assert colors(chart, ".//c:dPt//a:srgbClr") == ["FFDE00", "FF2E00", "FFDE00"]
    assert chart._chartSpace.xpath(".//c:invertIfNegative[@val='0']")
    assert chart.plots[0].data_labels.number_format == 'General" млн"'
    assert not chart.has_legend


def test_bar_chart_survives_special_characters_in_unit(tdir):
    deck = Deck(slides=[bar(unit="R&D <x>\x01", labels=["Центр"], series=[{"name": "R&D", "values": [42]}])])
    chart = chart_of(export(tdir, deck).slides[0])
    assert chart.plots[0].data_labels.number_format == 'General" R&D <x>"'


def test_bar_with_several_series_takes_the_palette_and_a_legend(tdir):
    deck = Deck(slides=[bar(series=[{"name": "2025 & <план>", "values": [1, -2, 3]},
                                    {"name": "2026", "values": [2, 3, 4]}])])
    chart = chart_of(export(tdir, deck).slides[0])
    assert colors(chart, ".//c:ser/c:spPr/a:solidFill/a:srgbClr") == ["FFDE00", "1E1E1E"]
    assert not chart._chartSpace.xpath(".//c:dPt")  # по знаку красится только одна серия
    assert chart.has_legend and chart.legend.position == XL_LEGEND_POSITION.TOP
    assert [s.name for s in chart.plots[0].series] == ["2025 & <план>", "2026"]


def test_bar_with_several_series_keeps_the_unit_on_the_axis(tdir):
    deck = Deck(slides=[bar(unit="A&B<", series=[{"name": "2025", "values": [1, 2, 3]},
                                                 {"name": "2026", "values": [2, 3, 4]}])])
    chart = chart_of(export(tdir, deck).slides[0])
    assert chart.value_axis.has_title and chart.value_axis.axis_title.text_frame.text == "A&B<"


def test_line_kind_draws_lines(tdir):
    chart = chart_of(export(tdir, Deck(slides=[bar(kind="line")])).slides[0])
    assert chart.chart_type == XL_CHART_TYPE.LINE_MARKERS


def test_dual_chart_two_charts_with_captions(tdir):
    side = {"labels": ["Q1", "Q2"], "series": [{"name": "Ряд", "values": [1, 2]}]}
    deck = Deck(slides=[{"type": "dual-chart", "title": "Т",
                         "left": side | {"caption": "Динамика & <рост>"},
                         "right": side | {"caption": "", "kind": "bar", "unit": "%"}}])
    slide = export(tdir, deck).slides[0]
    frames = sorted((sh for sh in slide.shapes if sh.has_chart), key=lambda sh: sh.left)
    assert len(frames) == 2 and all(inside(f) for f in frames)
    assert frames[0].chart.chart_type == XL_CHART_TYPE.LINE_MARKERS
    assert frames[1].chart.chart_type == XL_CHART_TYPE.COLUMN_CLUSTERED
    assert frames[1].left - (frames[0].left + frames[0].width) == 20 * MM
    [cap] = boxes(slide)  # у правой подписи нет — полосы под неё тоже
    assert cap.text_frame.text == "Динамика & <рост>" and cap.top < frames[0].top
    assert frames[1].top == int(ZONE[1] * MM)  # без подписи график начинается с верха зоны


def test_chart_with_zero_values_exports(tdir):
    deck = Deck(slides=[bar(labels=["A"], series=[{"name": "x", "values": [0]}])])
    assert len(export(tdir, deck).slides) == 1


def test_line_series_cycle_the_palette(tdir):
    deck = Deck(slides=[{"type": "chart", "title": "Итоги", "unit": "млн", "labels": ["Q1", "Q2"],
                         "series": [{"name": f"Ряд {j}", "values": [j, j + 1]} for j in range(4)]}])
    chart = chart_of(export(tdir, deck).slides[0])
    assert colors(chart, ".//c:ser/c:spPr/a:ln//a:srgbClr") == ["FFDE00", "1E1E1E", "00FFE0", "FFDE00"]
    assert chart.legend.position == XL_LEGEND_POSITION.TOP
    assert chart.value_axis.axis_title.text_frame.text == "млн"


def test_chart_takes_the_zone_or_shares_it_with_bullets(tdir):
    alone, shared = export(tdir, Deck(slides=[bar(), bar(bullets=["Рост"])])).slides
    frame = next(sh for sh in alone.shapes if sh.has_chart)
    assert (frame.left, frame.top, frame.width, frame.height) == tuple(int(v * MM) for v in ZONE)
    [text] = boxes(shared)
    frame = next(sh for sh in shared.shapes if sh.has_chart)
    assert text.left == int(ZONE[0] * MM) and text.width == int(ZONE[2] * MM * 38 / 100)
    assert frame.left == text.left + text.width + 12 * MM
    heading, bullet = text.text_frame.paragraphs
    assert heading.text == "Выводы" and not heading._p.xpath("./a:pPr/a:buChar")
    assert bullet._p.xpath("./a:pPr/a:buChar[@char='•']")


def test_agenda_is_a_numbered_list_in_the_zone(tdir):
    deck = Deck(slides=[{"type": "agenda", "items": ["Итоги", "Планы", "Риски"]}])
    slide = export(tdir, deck).slides[0]
    assert by_idx(slide)[0] == ["Повестка"]
    [box] = boxes(slide)
    assert inside(box) and [p.text for p in box.text_frame.paragraphs] == ["Итоги", "Планы", "Риски"]
    for p in box.text_frame.paragraphs:
        assert p._p.xpath("./a:pPr/a:buAutoNum[@type='arabicPeriod']")
        assert p._p.xpath("./a:pPr/a:buClr/a:srgbClr[@val='FFDE00']")


def test_text_bullets_take_the_whole_zone(tdir):
    slide = export(tdir, Deck(slides=[{"type": "text", "title": "Т", "bullets": ["раз", "два"]}])).slides[0]
    [box] = boxes(slide)
    assert (box.left, box.top, box.width, box.height) == tuple(int(v * MM) for v in ZONE)
    assert [p.text for p in box.text_frame.paragraphs] == ["раз", "два"]


def test_columns_split_the_zone(tdir):
    deck = Deck(slides=[{"type": "columns", "title": "Т", "blocks": [
        {"heading": "Первый", "text": "раз"}, {"heading": "", "text": "два"}, {"heading": "Третий", "text": "три"}]}])
    cols = sorted(boxes(export(tdir, deck).slides[0]), key=lambda b: b.left)
    assert len(cols) == 3 and all(inside(b) for b in cols)
    assert cols[1].left - (cols[0].left + cols[0].width) == 10 * MM  # gap
    first = cols[0].text_frame.paragraphs
    assert [p.text for p in first] == ["Первый", "раз"] and first[0].runs[0].font.bold
    assert [p.text for p in cols[1].text_frame.paragraphs] == ["два"]  # пустой заголовок не печатается


def test_columns_on_placeholders_drop_the_spare_one(tmp_path):
    tdir = make_pptx_template(tmp_path, mapping_edit=lambda m: m["slides"]["columns"].update(
        layout="Comparison", ph={"title": 0, "blocks": [1, 2, 3]}, body=None, gap=None))
    deck = Deck(slides=[{"type": "columns", "title": "Т", "blocks": [
        {"heading": "Первый", "text": "раз"}, {"heading": "Второй", "text": "два"}]}])
    slide = export(tdir, deck).slides[0]
    got = by_idx(slide)
    assert got[1] == ["Первый", "раз"] and got[2] == ["Второй", "два"]
    assert 3 not in got and 4 not in got  # третьего блока нет, idx 4 не сопоставлен
    assert boxes(slide) == []


def test_comparison_two_columns_with_headings(tdir):
    deck = Deck(slides=[{"type": "comparison", "title": "Т",
                         "left": {"heading": "Было", "bullets": ["а"]},
                         "right": {"heading": "Стало", "bullets": ["б", "в"]}}])
    left, right = sorted(boxes(export(tdir, deck).slides[0]), key=lambda b: b.left)
    assert [p.text for p in left.text_frame.paragraphs] == ["Было", "а"]
    assert [p.text for p in right.text_frame.paragraphs] == ["Стало", "б", "в"]
    assert left.width == right.width and right.left - (left.left + left.width) == 20 * MM


def test_quote_is_centered_on_a_layout_without_title(tdir):
    deck = Deck(slides=[{"type": "quote", "text": "Мысль", "author": "Автор"}])
    slide = export(tdir, deck).slides[0]
    [box] = boxes(slide)
    assert inside(box) and box.text_frame.vertical_anchor == MSO_ANCHOR.MIDDLE
    quote, author = box.text_frame.paragraphs
    assert quote.alignment == PP_ALIGN.CENTER and quote.runs[0].font.size == Pt(40)
    assert author.text == "Автор" and str(author.runs[0].font.color.rgb) == "666666"
    assert set(by_idx(slide)) <= {12}  # кроме номера слайда заполнителей нет


def test_text_fields_drop_control_characters(tdir):
    deck = Deck(slides=[
        {"type": "columns", "title": "Т", "blocks": [{"text": "раз\x1f\nдва"}, {"text": "три"}]},
        {"type": "quote", "text": "мысль\x0b"},
    ])
    prs = export(tdir, deck)
    xml = b"".join(etree.tostring(s._element) for s in prs.slides)
    assert b"_x001F_" not in xml and b"_x000B_" not in xml
    first = sorted(boxes(prs.slides[0]), key=lambda b: b.left)[0]
    assert first._element.xpath(".//a:br")  # \n — разрыв строки внутри абзаца


def test_metrics_columns_with_trend_colors(tdir):
    deck = Deck(slides=[{"type": "metrics", "title": "Т", "items": [
        {"value": "12,5 млн", "label": "выручка", "delta": "+12 %"},
        {"value": "3,4 %", "label": "доля", "delta": "–1 п. п."},
        {"value": "17", "label": "проектов\x1f"}]}])
    cols = sorted(boxes(export(tdir, deck).slides[0]), key=lambda b: b.left)
    assert len(cols) == 3 and all(inside(b) for b in cols)
    value, label, delta = cols[0].text_frame.paragraphs
    assert value.text == "12,5 млн" and value.runs[0].font.size == Pt(60) and value.runs[0].font.bold
    assert label.text == "выручка" and str(label.runs[0].font.color.rgb) == "666666"
    assert str(delta.runs[0].font.color.rgb) == "1E1E1E"  # рост — trend.up
    assert str(cols[1].text_frame.paragraphs[2].runs[0].font.color.rgb) == "FF2E00"  # «–» — падение
    assert [p.text for p in cols[2].text_frame.paragraphs] == ["17", "проектов"]  # без delta — два абзаца


def table_of(slide):
    return next(sh for sh in slide.shapes if sh.has_table)


def test_table_is_native_with_header_stripes_and_numeric_alignment(tdir):
    deck = Deck(slides=[{"type": "table", "title": "Т", "header": ["Регион", "Выручка"],
                         "rows": [["Центр", "42"], ["Север\nи запад", "17,5"], ["Юг", "28 %"]]}])
    frame = table_of(export(tdir, deck).slides[0])
    tbl = frame.table
    assert (len(tbl.rows), len(tbl.columns)) == (4, 2) and inside(frame)
    head = tbl.cell(0, 0)
    assert head.text == "Регион" and head.text_frame.paragraphs[0].runs[0].font.bold
    assert head._tc.xpath("./a:tcPr/a:solidFill/a:srgbClr[@val='DDDDDD']")
    assert tbl.cell(2, 0)._tc.xpath("./a:tcPr/a:solidFill/a:srgbClr[@val='F8F8F8']")  # вторая строка данных
    assert tbl.cell(1, 0)._tc.xpath("./a:tcPr/a:noFill")
    assert tbl.cell(1, 1).text_frame.paragraphs[0].alignment == PP_ALIGN.RIGHT  # числа вправо
    assert tbl.cell(1, 0).text_frame.paragraphs[0].alignment != PP_ALIGN.RIGHT
    assert tbl.cell(2, 0)._tc.xpath(".//a:br")
    # линии раньше заливки — иначе PowerPoint предложит восстановить файл
    tags = [etree.QName(e).localname for e in tbl.cell(0, 0)._tc.tcPr]
    assert tags[:4] == ["lnL", "lnR", "lnT", "lnB"] and tags[4] == "solidFill"
    assert not frame._element.xpath(".//a:tblPr[@bandRow='1' or @firstRow='1']")


def test_table_at_model_limits_lays_its_grid_inside_the_zone(tdir):
    """Номинальная сетка — в зоне. Перенос текста PowerPoint считает сам и может
    нарастить строки; это проверяется ручной приёмкой таблицы на пределах."""
    from typst_deck.model import CELL_MAX

    cell = "x" * CELL_MAX
    deck = Deck(slides=[{"type": "table", "title": "Т", "header": ["h" * 40] * 6, "rows": [[cell] * 6] * 8}])
    frame = table_of(export(tdir, deck).slides[0])
    assert inside(frame) and frame.height <= int(ZONE[3] * MM)


def test_short_table_rows_are_capped(tdir):
    deck = Deck(slides=[{"type": "table", "title": "Т", "header": ["a", "b"], "rows": [["1", "2"]]}])
    frame = table_of(export(tdir, deck).slides[0])
    assert frame.table.rows[0].height == 15 * MM and frame.height == 30 * MM


def gallery(n, caption="Подпись"):
    return Deck(slides=[{"type": "gallery", "title": "Г",
                         "images": [{"image": f"p{i}.png", "caption": caption} for i in range(n)]}])


@pytest.mark.parametrize("n,rows,cols", [(2, 1, 2), (3, 1, 3), (4, 2, 2)])
def test_gallery_grid(tdir, tmp_path, n, rows, cols):
    (tmp_path / "a").mkdir()
    for i in range(n):
        png(tmp_path / "a" / f"p{i}.png")
    slide = export(tdir, gallery(n), assets=tmp_path / "a").slides[0]
    pics = pictures(slide)
    assert len(pics) == n and all(inside(p) for p in pics)
    assert len({p.left for p in pics}) == cols and len({p.top for p in pics}) == rows
    assert all(abs(p.width / p.height - 2) < 0.01 for p in pics)
    caps = boxes(slide)
    assert len(caps) == n and all(inside(b) for b in caps)


def test_gallery_missing_picture_keeps_the_grid(tdir, tmp_path):
    (tmp_path / "a").mkdir()
    png(tmp_path / "a" / "p0.png")
    warnings: list[str] = []
    slide = export(tdir, gallery(2), assets=tmp_path / "a", warnings=warnings).slides[0]
    [pic] = pictures(slide)
    assert pic.left < int((ZONE[0] + ZONE[2] / 2) * MM)  # осталась левая, правая клетка пустая
    assert any("p1.png не найдена" in w for w in warnings)


def video_deck(**kw):
    return Deck(slides=[{"type": "video", "title": "В", "caption": "Ролик", "url": "https://example.com/v"} | kw])


def test_video_poster_carries_the_link(tdir, tmp_path):
    (tmp_path / "a").mkdir()
    png(tmp_path / "a" / "cover.png", 320, 240)
    url = "https://vk.com/video-1_2?list=a&t=3"
    slide = export(tdir, video_deck(url=url, poster="cover.png"), assets=tmp_path / "a").slides[0]
    [pic] = pictures(slide)
    assert pic.click_action.hyperlink.address == url and abs(pic.width / pic.height - 4 / 3) < 0.01
    [cap] = boxes(slide)
    assert cap.text_frame.text == "Ролик" and cap.top >= pic.top + pic.height


def test_missing_poster_falls_back_to_the_placeholder(tdir, tmp_path):
    warnings: list[str] = []
    slide = export(tdir, video_deck(poster="gone.png"), assets=tmp_path, warnings=warnings).slides[0]
    [pic] = pictures(slide)
    assert abs(pic.width / pic.height - 16 / 9) < 0.01 and any("gone.png не найдена" in w for w in warnings)


def test_link_video_is_a_poster_with_a_hyperlink(tdir):
    url = "https://example.com/watch?v=1&t=2"
    slide = export(tdir, video_deck(url=url)).slides[0]
    [pic] = pictures(slide)
    assert pic.click_action.hyperlink.address == url
    assert abs(pic.width / pic.height - 16 / 9) < 0.01  # заглушка 16:9
    assert not [sh for sh in slide.shapes if sh.shape_type == MSO_SHAPE_TYPE.MEDIA]


def test_placeholder_poster_takes_the_mapping_colors(tdir):
    slide = export(tdir, video_deck(url="https://example.com")).slides[0]
    [pic] = pictures(slide)
    with Image.open(io.BytesIO(pic.image.blob)) as im:
        assert im.getpixel((5, 5)) == (0xDD, 0xDD, 0xDD)  # table.header_fill
        assert im.getpixel((im.width // 2, im.height // 2)) == (0xFF, 0xDE, 0x00)  # «▶» — text.bullet


def test_notes_go_to_speaker_notes(tdir):
    deck = Deck(slides=[{"type": "final", "notes": "Сказать спасибо"}])
    assert export(tdir, deck).slides[0].notes_slide.notes_text_frame.text == "Сказать спасибо"


def test_control_characters_dropped_and_newline_breaks_the_line(tdir):
    deck = Deck(slides=[{"type": "title", "title": "\x0bИтоги\x1f", "subtitle": "первая\nвторая"},
                        {"type": "text", "title": "T", "bullets": ["тезис\nпродолжение"]}])
    prs = export(tdir, deck)
    assert by_idx(prs.slides[0])[0] == ["Итоги"]
    sub = next(sh for sh in prs.slides[0].placeholders if sh.placeholder_format.idx == 1)
    assert len(sub.text_frame.paragraphs) == 1 and sub._element.xpath(".//a:br")
    [text] = boxes(prs.slides[1])
    assert len(text.text_frame.paragraphs) == 1 and text._element.xpath(".//a:br")


def test_text_at_global_limits_exports(tdir):
    """Каждый тип на глобальных пределах модели собирается и открывается заново."""
    from typst_deck.model import SLIDE_TYPES as ALL

    deck = Deck(slides=[filled_slide(t, {"title": {"max": 120, "lines": 2}} if t != "quote" else {})
                        for t in ALL])
    assert len(export(tdir, deck).slides) == len(ALL)


def latin_faces(element) -> list[str]:
    return [e.get("typeface") for e in element.iter(qn("a:latin"))]


def theme_faces(prs) -> list[str]:
    theme = prs.slide_masters[0].part.part_related_by(RT.THEME)
    return latin_faces(etree.fromstring(theme.blob))


def with_helvetica_layout(tdir):
    """Явная гарнитура в макете «Title Only» — как у мастеров, где шрифт прописан мимо темы."""
    prs = Presentation(str(tdir / "master.pptx"))
    run = prs.slide_layouts.get_by_name("Title Only").placeholders[0].text_frame.paragraphs[0].add_run()
    run.font.name = "Helvetica"
    prs.save(str(tdir / "master.pptx"))


def test_template_font_replaces_explicit_latin_faces(tmp_path):
    tdir = make_pptx_template(tmp_path, mapping_edit=lambda m: m.update(font="Arial"))
    with_helvetica_layout(tdir)
    prs = export(tdir, SAMPLE_DECK)
    assert set(theme_faces(prs)) == {"Arial"}  # основной текст, заголовки, умолчания новых фигур
    master = prs.slide_masters[0]
    assert set(latin_faces(master._element)) == {"+mj-lt", "+mn-lt"}  # ссылки на тему не трогаются
    assert latin_faces(master.slide_layouts.get_by_name("Title Only")._element) == ["Arial"]


def test_template_font_reaches_speaker_notes(tmp_path):
    """Заметки набираются по теме мастера заметок: у Lemana Tech она своя, с Helvetica."""
    tdir = make_pptx_template(tmp_path, mapping_edit=lambda m: m.update(font="Arial"))
    prs = export(tdir, Deck(slides=[{"type": "final", "notes": "Сказать спасибо"}]))
    notes_master = prs.part.part_related_by(RT.NOTES_MASTER)
    assert set(latin_faces(etree.fromstring(notes_master.part_related_by(RT.THEME).blob))) == {"Arial"}


def test_without_font_the_master_keeps_its_faces(tmp_path):
    tdir = make_pptx_template(tmp_path)
    with_helvetica_layout(tdir)
    prs = export(tdir, SAMPLE_DECK)
    assert set(theme_faces(prs)) == {"Calibri"}
    assert latin_faces(prs.slide_masters[0].slide_layouts.get_by_name("Title Only")._element) == ["Helvetica"]


# --- сопоставление

def check(tdir):
    open_master(tdir, load_mapping(tdir, "pptx.json"))


def test_good_mapping_passes(tmp_path):
    check(make_pptx_template(tmp_path))


BROKEN = {
    "master-outside": (lambda m: m.update(master="../master.pptx"), "внутрь шаблона"),
    "master-missing": (lambda m: m.update(master="nope.pptx"), "нет файла nope.pptx"),
    "unknown-key": (lambda m: m.update(colour="FFFFFF"), "colour"),
    "missing-type": (lambda m: m["slides"].pop("final"), "нет типа final"),
    "missing-ph-field": (lambda m: m["slides"]["title"]["ph"].pop("department"), "нет поля department"),
    "extra-ph-field": (lambda m: m["slides"]["image"]["ph"].update(caption=1), "caption не кладётся"),
    "no-body": (lambda m: m["slides"]["image"].pop("body"), "нет body"),
    "body-on-title": (lambda m: m["slides"]["title"].update(body=[0, 0, 10, 10]), "body у этого типа не бывает"),
    "bad-color": (lambda m: m["chart"].update(negative="red"), "chart.negative"),
    "empty-font": (lambda m: m.update(font=""), "font"),
    "no-layout": (lambda m: m["slides"]["section"].update(layout="Нет такого"), "нет макета «Нет такого»"),
    "no-idx": (lambda m: m["slides"]["title"]["ph"].update(author=42), "нет заполнителя idx 42"),
    "slide-number-idx": (lambda m: m["slides"]["title"]["ph"].update(subtitle=12), "нет заполнителя idx 12"),
    "image-no-room-for-text": (lambda m: m["slides"]["image"].update(
        image_width={"small": 90, "medium": 90, "large": 90}, gap=100),
        "при image_width 90 % и gap 100 мм колонке тезисов не остаётся места"),
    "chart-no-room-for-chart": (lambda m: m["slides"]["chart"].update(text_width=90, gap=70),
        "при text_width 90 % и gap 70 мм графику не остаётся места"),
    "image-body-too-short": (lambda m: m["slides"]["image"].update(body=[20, 80, 600, 10]),
        "высота 10 мм не больше полосы подписи 15 мм"),
    "zone-outside": (lambda m: m["slides"]["chart"].update(body=[600, 80, 200, 260]), "выходит за слайд"),
    "missing-new-type": (lambda m: m["slides"].pop("metrics"), "нет типа metrics"),
    "old-type": (lambda m: m["slides"].update({"bar-chart": m["slides"]["chart"]}), "bar-chart"),
    "no-trend": (lambda m: m.pop("trend"), "trend"),
    "no-table-style": (lambda m: m.pop("table"), "table"),
    "no-quote-size": (lambda m: m["text"].pop("quote"), "text.quote"),
    "quote-title-ph": (lambda m: m["slides"]["quote"]["ph"].update(title=0), "title не кладётся"),
    "no-gap": (lambda m: m["slides"]["metrics"].pop("gap"), "slides.metrics: нет gap"),
    "blocks-not-columns": (lambda m: m["slides"]["text"]["ph"].update(blocks=[1, 2, 3]), "blocks не кладётся"),
    "blocks-too-few": (lambda m: m["slides"]["columns"].update(
        layout="Comparison", ph={"title": 0, "blocks": [1, 2]}, body=None, gap=None),
        "нужно не меньше 3 заполнителей"),
    "blocks-and-zone": (lambda m: m["slides"]["columns"].update(
        layout="Comparison", ph={"title": 0, "blocks": [1, 2, 3]}),
        "slides.columns: body у этого типа не бывает"),
    "blocks-idx-missing": (lambda m: m["slides"]["columns"].update(
        layout="Comparison", ph={"title": 0, "blocks": [1, 2, 9]}, body=None, gap=None),
        "нет заполнителя idx 9"),
    "title-as-list": (lambda m: m["slides"]["text"]["ph"].update(title=[0]), "ожидается один idx"),
    "blocks-over-title": (lambda m: m["slides"]["columns"].update(
        layout="Comparison", ph={"title": 0, "blocks": [0, 1, 2]}, body=None, gap=None),
        "blocks: idx 0 занят другим полем"),
    "blocks-duplicate": (lambda m: m["slides"]["columns"].update(
        layout="Comparison", ph={"title": 0, "blocks": [1, 1, 2]}, body=None, gap=None),
        "blocks: idx 1 повторяется"),
    "gallery-no-room-across": (lambda m: m["slides"]["gallery"].update(gap=400),
        "slides.gallery: при gap 400 мм на колонку не остаётся места"),
    "metrics-no-room-across": (lambda m: m["slides"]["metrics"].update(gap=400),
        "slides.metrics: при gap 400 мм на колонку не остаётся места"),
    "columns-no-room-across": (lambda m: m["slides"]["columns"].update(gap=400), "slides.columns: при gap 400"),
    "comparison-no-room-across": (lambda m: m["slides"]["comparison"].update(gap=700),
        "slides.comparison: при gap 700"),
    "dual-chart-no-room-across": (lambda m: m["slides"]["dual-chart"].update(gap=700),
        "slides.dual-chart: при gap 700"),
    "gallery-too-short": (lambda m: m["slides"]["gallery"].update(body=[20, 80, 600, 30]),
        "slides.gallery.body: высота 30 мм мала для двух рядов с подписями"),
    "video-too-short": (lambda m: m["slides"]["video"].update(body=[20, 80, 600, 15]),
        "slides.video.body: высота 15 мм не больше полосы подписи 15 мм"),
    "dual-chart-too-short": (lambda m: m["slides"]["dual-chart"].update(body=[20, 80, 600, 12]),
        "slides.dual-chart.body: высота 12 мм не больше полосы подписи 15 мм"),
}


def test_columns_on_placeholders_is_a_good_mapping(tmp_path):
    check(make_pptx_template(tmp_path, mapping_edit=lambda m: m["slides"]["columns"].update(
        layout="Comparison", ph={"title": 0, "blocks": [1, 2, 3]}, body=None, gap=None)))


@pytest.mark.parametrize("case", BROKEN)
def test_broken_mapping_is_named(case, tmp_path):
    edit, text = BROKEN[case]
    d = make_pptx_template(tmp_path, mapping_edit=edit)
    with pytest.raises(MappingError, match=re.escape(text)):
        check(d)


def test_unreadable_mapping_is_named(tmp_path):
    d = make_pptx_template(tmp_path)
    (d / "pptx.json").write_text("{")
    with pytest.raises(MappingError, match="не читается"):
        load_mapping(d, "pptx.json")


def test_mapping_path_stays_inside_the_template(tmp_path):
    d = make_pptx_template(tmp_path)
    with pytest.raises(MappingError, match="внутрь шаблона"):
        load_mapping(d, "../pptx.json")
    with pytest.raises(MappingError, match="нет файла missing.json"):
        load_mapping(d, "missing.json")


def test_broken_master_file_is_named(tmp_path):
    d = make_pptx_template(tmp_path)
    (d / "master.pptx").write_bytes(b"not a zip")
    with pytest.raises(MappingError, match="не открывается"):
        check(d)


def _with_rule(tdir, rule):
    import json
    p = tdir / "template.json"
    meta = json.loads(p.read_text())
    meta["charts"] = rule
    p.write_text(json.dumps(meta, ensure_ascii=False))


def test_bars_turn_horizontal_by_the_template_rule(tdir):
    """Правило шаблона (charts.columns) — то же, что у PDF: длинные подписи — линейчатая
    диаграмма, первая категория сверху; короткие — гистограмма, как раньше."""
    _with_rule(tdir, {"columns": {"full": [[6, 8]], "bullets": [[6, 4]], "dual": [[6, 3]]}})
    long_ = bar(labels=["Стройматериалы", "Инструмент", "Сантехника"])
    pres = export(tdir, Deck(slides=[long_, bar(), bar(bullets=["вывод"], labels=list("ABCDE"), series=[{"name": "с", "values": [1, 2, 3, 4, 5]}])]))
    hbar, short, beside = (chart_of(s) for s in pres.slides)
    assert hbar.chart_type == XL_CHART_TYPE.BAR_CLUSTERED and hbar.category_axis.reverse_order
    assert short.chart_type == XL_CHART_TYPE.COLUMN_CLUSTERED                 # «Центр» — 5 знаков, 3 ≤ 8
    assert beside.chart_type == XL_CHART_TYPE.BAR_CLUSTERED                   # с выводами: 5 > 4


def test_dual_chart_side_uses_the_dual_layout(tdir):
    _with_rule(tdir, {"columns": {"full": [[30, 24]], "dual": [[4, 2]]}})
    side = {"kind": "bar", "labels": ["Q1", "Q2", "Q3"], "series": [{"name": "Ряд", "values": [1, 2, 3]}]}
    slide = export(tdir, Deck(slides=[{"type": "dual-chart", "title": "Т", "left": side, "right": side}])).slides[0]
    assert all(sh.chart.chart_type == XL_CHART_TYPE.BAR_CLUSTERED for sh in slide.shapes if sh.has_chart)


def test_broken_rule_keeps_columns_vertical_with_a_warning(tdir):
    _with_rule(tdir, {"columns": {"wide": []}})
    warnings = []
    pres = export(tdir, Deck(slides=[bar(labels=["Стройматериалы"] * 3)]), warnings=warnings)
    assert chart_of(pres.slides[0]).chart_type == XL_CHART_TYPE.COLUMN_CLUSTERED
    assert any("charts.columns" in w for w in warnings)
