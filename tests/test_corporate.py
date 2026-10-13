"""Общие тесты корпоративных шаблонов: те же проверки для каждого шаблона из
TYPST_DECK_TEST_TEMPLATES. Раньше они были прибиты к templates/lmru-hr; теперь
корпоративный шаблон живёт в репозитории своего канона (docs/template-contract.md).
На этом сервере: TYPST_DECK_TEST_TEMPLATES=~/corp-templates uv run pytest
"""
import io
import json

import pytest
from pptx import Presentation

from tests.conftest import CORP_PARAMS, NO_CORP, requires_typst
from tests.test_charts import signed_deck
from typst_deck.contract import check_template, filled_slide, SAMPLE_DECK
from typst_deck.model import SLIDE_TYPES, Deck
from typst_deck.pptx import export_pptx
from typst_deck.render import compile_deck
from typst_deck.slots import check_slide, clamp

MM = 2.83465  # мм в пунктах: зоны заданы в мм, positions отдаёт пункты
TYPE_PARAMS = ([pytest.param(t, s, id=f"{t.id}-{s}") for p in CORP_PARAMS for t in p.values if t is not None
                for s in SLIDE_TYPES]
               or [pytest.param(None, None, id="no-corp", marks=pytest.mark.skip(reason=NO_CORP))])
# Низ тела content-slide канона: поле 24 мм + зона заголовка 50 мм + отбивка 6 мм +
# межабзацный отступ канона перед телом (1,2em при 28pt ≈ 11,9 мм) + тело 250 мм —
# верх тела у всех полей на 91,9 мм (по positions). Тело обрезано clip — всё, что
# ниже, теряется молча; граница по полю страницы (357 мм) это пропускала.
PAGE_BOTTOM_MM = 91.9 + 250


def _model_maxima(stype: str) -> dict:
    """Глобальные пределы модели для каждого поля типа (и подполей массивов объектов):
    слайд «на пределах» — это и необъявленные поля по максимуму, а не их короткие
    значения по умолчанию из filled_slide."""
    from typst_deck.slots import DEFS, SCHEMA, _ceiling, _resolve

    props, defs = DEFS[stype]["properties"], SCHEMA["$defs"]
    out: dict = {}
    for f, prop in props.items():
        if f in ("url", "poster", "image"):
            continue  # ссылка и имена ассетов — форма, а не текст: образцом по длине их не заполнить
        lim = {k: c for k in ("max", "count") if (c := _ceiling(prop, k)) is not None}
        if lim:
            out[f] = lim
        sub = prop.get("items", prop) if prop.get("type") == "array" else prop
        ref = sub.get("$ref")
        target = defs.get(ref.rsplit("/", 1)[1], {}) if ref else sub
        for g in target.get("properties", {}):
            sp = _resolve(props, defs, f"{f}.{g}")
            sl = {k: c for k in ("max", "count") if sp and (c := _ceiling(sp, k)) is not None}
            if sl:
                out[f"{f}.{g}"] = sl
    return out


@requires_typst
@pytest.mark.parametrize("t", CORP_PARAMS)
def test_corp_template_passes_the_full_contract(t, tmp_path):
    assert check_template(t.path, "typst", tmp_path).errors == []


@pytest.mark.parametrize("t", CORP_PARAMS)
def test_corp_declaration_survives_clamp_unchanged(t):
    raw = json.loads((t.path / "template.json").read_text()).get("slots") or {}
    assert clamp(raw, t.id) == raw


@pytest.mark.parametrize("t", CORP_PARAMS)
def test_corp_rejects_the_2026_09_09_incident(t):
    """Переводы строки в подзаголовке титульного слайда сверх объявленного — отказ до рендера."""
    lines = t.slots.get("title", {}).get("subtitle", {}).get("lines")
    if lines is None:
        pytest.skip(f"{t.id}: у подзаголовка титульного слайда не объявлено lines")
    bad = Deck(slides=[{"type": "title", "title": "Итоги",
                        "subtitle": "Лемана" + "\n" * lines + "хвост"}]).slides[0]
    assert check_slide(bad, t.slots, n=1, template_name=t.name)
    ok = Deck(slides=[{"type": "title", "title": "Итоги",
                       "subtitle": "\n".join(["Лемана"] * lines)}]).slides[0]
    assert check_slide(ok, t.slots, n=1, template_name=t.name) == []


@pytest.mark.parametrize("t", CORP_PARAMS)
def test_corp_every_declared_limit_is_enforced(t):
    for stype, fields in t.slots.items():
        for fld, lim in fields.items():
            for key, value in lim.items():
                over = dict(lim)
                over[key] = value + 1
                try:
                    slide = Deck(slides=[filled_slide(stype, {fld: over})]).slides[0]
                except ValueError:
                    continue  # сверх глобального потолка model.py — его держит валидация модели
                assert check_slide(slide, t.slots, n=1, template_name=t.name), f"{stype}.{fld}.{key}"


@requires_typst
@pytest.mark.parametrize("t", CORP_PARAMS)
def test_corp_declared_text_fits_its_measured_zone(t, tmp_path):
    """Текст ровно по объявленным пределам не выше зоны, в которой его меряли (measure.json).
    Обрезку внутри слота правило «одна страница» не видит, а этот тест видит."""
    meta = json.loads((t.path / "template.json").read_text())
    if not meta.get("measure"):
        pytest.skip(f"{t.id}: нет measure")
    cfg = json.loads((t.path / meta["measure"]).read_text())
    for i, p in enumerate(cfg["probes"]):
        lim = t.slots.get(p["slide"], {}).get(p["field"])
        if not lim or "." in p["field"]:
            # вложенное поле (images.caption) помечено маркером поля верхнего уровня
            # вместе с соседями (картинкой) — его высоту отдельно не достать;
            # вложенные пределы держит test_corp_type_at_declared_limits_stays_on_the_page
            continue
        slide = filled_slide(p["slide"], {p["field"]: lim}, sample=cfg.get("sample", "Итоги квартала "))
        r = compile_deck(Deck(slides=[slide]), tmp_path / f"p{i}", t.path, fonts_dir=t.fonts)
        h = next(x["h"] for x in r.positions if x["field"] == p["field"])
        zone = float(p["height"].removesuffix("mm")) * MM
        assert h <= zone, f"{p['slide']}.{p['field']}: {h:.1f}pt при зоне {zone:.1f}pt"


@requires_typst
@pytest.mark.parametrize("t", CORP_PARAMS)
def test_corp_signed_charts_keep_one_page_per_slide(t, deck_dir):
    assert len(compile_deck(signed_deck(), deck_dir, t.path, fonts_dir=t.fonts).svgs) == 2


@pytest.mark.parametrize("t", CORP_PARAMS)
def test_corp_pptx_exports_the_sample_and_declared_limits(t, tmp_path):
    if t.pptx is None:
        pytest.skip(f"{t.id}: нет PowerPoint-мастера")
    declared = Deck(slides=[filled_slide(s, lim) for s, lim in t.slots.items()])
    for deck in (SAMPLE_DECK, declared):
        data = export_pptx(deck, t.path, t.pptx, tmp_path, "typst")
        assert len(Presentation(io.BytesIO(data)).slides) == len(deck.slides)


def test_corp_lemana_tech_has_a_master():
    """Страховка от тихого отката: у канона Lemana Tech мастер обязан быть подключён."""
    lemana = [t for t in (p.values[0] for p in CORP_PARAMS) if t is not None and t.id == "lemana-tech"]
    if not lemana:
        pytest.skip("lemana-tech не подключён в TYPST_DECK_TEST_TEMPLATES")
    assert lemana[0].pptx == "typst-deck/pptx.json"


@requires_typst
@pytest.mark.parametrize("t,stype", TYPE_PARAMS)
def test_corp_type_compiles_on_one_page(t, stype, tmp_path):
    """Минимальный слайд каждого типа — одна страница: функция контракта есть и не переносит."""
    from typst_deck.contract import write_sample_assets

    write_sample_assets(tmp_path)
    slide = next(s for s in SAMPLE_DECK.slides if s.type == stype)
    r = compile_deck(Deck(slides=[slide]), tmp_path, t.path, fonts_dir=t.fonts)
    assert len(r.svgs) == 1


@requires_typst
@pytest.mark.parametrize("t,stype", TYPE_PARAMS)
def test_corp_type_at_declared_limits_stays_on_the_page(t, stype, tmp_path):
    """Слайд по объявленным пределам шаблона (а без объявления — по глобальным модели):
    ни одно поле не уходит ниже поля страницы — обрезку clip правило «одна страница» не видит."""
    from typst_deck.contract import write_sample_assets

    write_sample_assets(tmp_path)
    # объявленное шаблоном поверх глобальных пределов модели
    slide = filled_slide(stype, _model_maxima(stype) | t.slots.get(stype, {}))
    r = compile_deck(Deck(slides=[slide]), tmp_path, t.path, fonts_dir=t.fonts)
    assert len(r.svgs) == 1
    worst = max((p["y"] + p["h"] for p in r.positions), default=0)
    assert worst <= PAGE_BOTTOM_MM * MM, f"{stype}: низ {worst / MM:.1f} мм при поле {PAGE_BOTTOM_MM} мм"


DEGENERATE = {
    "одна категория": {"labels": ["Центр"], "series": [{"name": "млн", "values": [42]}]},
    "все нули": {"labels": ["A", "B", "C"], "series": [{"name": "x", "values": [0, 0, 0]}]},
    "только минус": {"labels": ["A", "B"], "series": [{"name": "x", "values": [-3, -7]}]},
    "три ряда": {"labels": ["A", "B", "C"],
                 "series": [{"name": f"Ряд {k}", "values": [k, -k, 2 * k]} for k in range(1, 4)]},
}


@requires_typst
@pytest.mark.parametrize("t", CORP_PARAMS)
@pytest.mark.parametrize("kind", ["bar", "line"])
@pytest.mark.parametrize("case", DEGENERATE)
def test_corp_degenerate_charts_stay_on_one_page(t, kind, case, tmp_path):
    data = DEGENERATE[case]
    if kind == "line" and len(data["labels"]) < 2:
        pytest.skip("линии нужны две точки — модель это не пропустит")
    deck = Deck(slides=[{"type": "chart", "title": "Итоги", "kind": kind, "unit": "млн", **data},
                        {"type": "dual-chart", "title": "Два", "left": {"kind": kind, **data},
                         "right": {"kind": kind, **data}}])
    r = compile_deck(deck, tmp_path, t.path, fonts_dir=t.fonts)
    assert len(r.svgs) == 2
    assert max(p["y"] + p["h"] for p in r.positions) <= PAGE_BOTTOM_MM * MM


def _yellow_runs(png_path, y_frac):
    """Отрезки жёлтого (#FFDE00) в строке картинки на высоте y_frac — ширины столбцов."""
    from PIL import Image

    im = Image.open(png_path).convert("RGB")
    w, h = im.size
    y = int(h * y_frac)
    runs, start = [], None
    for x in range(w):
        r, g, b = im.getpixel((x, y))
        yellow = r > 230 and 200 < g < 240 and b < 60
        if yellow and start is None:
            start = x
        if not yellow and start is not None:
            runs.append(x - start)
            start = None
    if start is not None:
        runs.append(w - start)
    return runs


@requires_typst
@pytest.mark.parametrize("t", CORP_PARAMS)
def test_corp_chart_with_bullets_keeps_the_last_bar_whole(t, tmp_path):
    """Колонка выводов и график не шире тела: иначе clip срезает последний столбец справа
    (было: 38 % + 62 % + зазор 12 мм — на 12 мм шире тела). У столбцов нет маркеров поля,
    поэтому проверка — по картинке: при равных значениях все столбцы одной ширины.
    Десять категорий, а не двенадцать: у шаблона с правилом ориентации (charts.columns)
    двенадцать столбцов рядом с выводами уже горизонтальны — их стережёт тест ниже."""
    import subprocess

    deck = Deck(slides=[{"type": "chart", "title": "Итоги", "kind": "bar", "labels": [f"М{i}" for i in range(1, 11)],
                         "series": [{"name": "x", "values": [5] * 10}], "bullets": ["Вывод"]}])
    compile_deck(deck, tmp_path, t.path, fonts_dir=t.fonts)
    png = tmp_path / "p.png"
    subprocess.run(["typst", "compile", "--root", str(tmp_path / "build"), "--font-path", str(t.fonts),
                    "--format", "png", "--ppi", "40", str(tmp_path / "build" / "deck.typ"), str(png)], check=True,
                   capture_output=True)
    runs = [r for r in _yellow_runs(png, 0.6) if r > 3]
    assert len(runs) == 10, runs
    assert max(runs) - min(runs) <= 2, f"столбцы разной ширины — последний срезан: {runs}"


def _yellow_column_runs(png_path, x_frac):
    """Отрезки жёлтого в столбце картинки на ширине x_frac — толщины горизонтальных столбцов."""
    from PIL import Image

    im = Image.open(png_path).convert("RGB").transpose(Image.Transpose.TRANSPOSE)
    tmp = png_path.with_suffix(".t.png")
    im.save(tmp)
    return _yellow_runs(tmp, x_frac)


@requires_typst
@pytest.mark.parametrize("t", CORP_PARAMS)
def test_corp_bars_turn_horizontal_by_the_template_rule(t, tmp_path):
    """Шаблон с правилом charts.columns: 24 подписи по 30 знаков под вертикальными
    столбцами не помещаются (налезали друг на друга) — столбцы горизонтальны, по одному
    на категорию. Шаблон без правила — пропуск: у него столбцы всегда вертикальны."""
    import subprocess

    from typst_deck.charts import horizontal, load_rule

    src = "Стройматериалы Инструмент Сантехника Электрика Напольные покрытия Сад и огород " * 3
    labels = [src[(i * 7) % 60:(i * 7) % 60 + 30].strip() for i in range(24)]
    if not horizontal(labels, "full", load_rule(t.path)):
        pytest.skip("у шаблона нет правила ориентации, по которому эти столбцы горизонтальны")
    deck = Deck(slides=[{"type": "chart", "title": "Итоги", "kind": "bar", "labels": labels,
                         "series": [{"name": "x", "values": [5] * 24}]}])
    compile_deck(deck, tmp_path, t.path, fonts_dir=t.fonts)
    png = tmp_path / "p.png"
    subprocess.run(["typst", "compile", "--root", str(tmp_path / "build"), "--font-path", str(t.fonts),
                    "--format", "png", "--ppi", "40", str(tmp_path / "build" / "deck.typ"), str(png)], check=True,
                   capture_output=True)
    runs = [r for r in _yellow_column_runs(png, 0.6) if r > 1]
    assert len(runs) == 24, runs
