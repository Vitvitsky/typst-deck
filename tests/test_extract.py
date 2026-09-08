import io

import pytest

from typst_deck.extract import UnsupportedFile, extract


def test_txt_md_csv():
    assert extract("a.txt", "привет".encode()).text == "привет"
    assert extract("a.md", b"# H").text == "# H"
    t = extract("a.csv", "a,b\n1,2\n".encode()).text
    assert "| a | b |" in t and "| 1 | 2 |" in t


def test_docx():
    import docx
    d = docx.Document()
    d.add_paragraph("Первый абзац")
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "к1", "к2"
    buf = io.BytesIO()
    d.save(buf)
    out = extract("doc.docx", buf.getvalue()).text
    assert "Первый абзац" in out and "| к1 | к2 |" in out


def test_xlsx_limits_rows():
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Лист"
    for i in range(300):
        ws.append([i, i * 2])
    buf = io.BytesIO()
    wb.save(buf)
    out = extract("t.xlsx", buf.getvalue()).text
    assert "## Лист" in out and "| 199 | 398 |" in out and "| 250 | 500 |" not in out


def test_truncate():
    e = extract("a.txt", b"x" * 100, limit=10)
    assert e.truncated and len(e.text) < 100 and "усечено" in e.text


def test_images_and_unknown(png_bytes):
    assert extract("p.png", png_bytes).kind == "image"
    with pytest.raises(UnsupportedFile):
        extract("x.exe", b"MZ")
