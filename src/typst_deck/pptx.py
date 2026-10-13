"""Экспорт колоды в редактируемый PPTX на PowerPoint-мастере шаблона.

Спека: docs/superpowers/specs/2026-09-11-pptx-export-design.md. Сопоставление
(pptx.json в шаблоне, путь — поле pptx в template.json) говорит, на какой макет
мастера ложится каждый тип слайда, в какой заполнитель — каждое поле и где зона
тела. Продукт ничего корпоративного не знает: мастер, макеты и цвета приходят
из шаблона. Контракт (правило 7) и сам экспорт проверяют сопоставление одними
и теми же load_mapping и open_master.
"""
from __future__ import annotations

import copy
import io
import json
import re
import tempfile
import uuid
from dataclasses import dataclass
from pathlib import Path
from typing import Annotated, Callable, Literal

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.util import Emu, Mm, Pt
from pydantic import BaseModel, ConfigDict, Field, ValidationError

from .model import SLIDE_TYPES, Deck
from .render import CompileError, inside_template, rasterize_svg

MM = 36000  # EMU в миллиметре
# Поля, которые обязаны лежать в заполнителях, — в порядке модели: он же задаёт
# порядок абзацев, когда несколько полей делят один заполнитель. Типам, которых
# здесь нет, нужен только заголовок. Остальное идёт в зону body, notes — в заметки.
PH_FIELDS = {
    "title": ("title", "subtitle", "author", "department"),
    "section": ("title", "subtitle"),
    "final": ("title", "contact"),
    "quote": (),
}
# Ключи зоны тела: у типа из этого словаря они обязательны, у прочих запрещены.
BODY_KEYS = {
    "agenda": ("body",), "text": ("body",), "quote": ("body",),
    "table": ("body",), "video": ("body",),
    "image": ("body", "image_width", "gap"),
    "columns": ("body", "gap"),  # только без ph.blocks — см. _body_keys
    "comparison": ("body", "gap"), "dual-chart": ("body", "gap"),
    "metrics": ("body", "gap"), "gallery": ("body", "gap"),
    "chart": ("body", "text_width", "gap"),
}
OPTIONAL_KEYS = {"chart": ("heading",)}
BLOCKS = 3  # столько блоков допускает ColumnsSlide.blocks (max_length в model.py)

CAPTION_STRIP = 15   # мм под подпись к картинке
CAPTION_SCALE = 0.7  # кегль подписи от кегля тезисов
# Управляющие символы, которых не бывает в тексте слайда; \t и \n остаются.
# python-pptx на них не падает, а пишет в текст видимый мусор вида _x001F_.
CONTROL = re.compile(r"[\x00-\x08\x0b-\x1f]")

Hex = Annotated[str, Field(pattern=r"^[0-9A-Fa-f]{6}$")]
Percent = Annotated[float, Field(ge=1, le=90)]


class MappingError(ValueError):
    """Сопоставление или мастер шаблона не годятся; текст — для человека."""


class _Strict(BaseModel):
    # опечатка в ключе не должна тихо выключать настройку
    model_config = ConfigDict(extra="forbid")


class TextStyle(_Strict):
    size: float = Field(gt=0)    # кегль тезисов, пт
    bullet: Hex                  # цвет маркера тезисов и номеров повестки
    caption: Hex                 # цвет подписей: картинки, автора цитаты, метрики
    quote: float = Field(gt=0)   # кегль цитаты, пт
    metric: float = Field(gt=0)  # кегль крупного числа метрики, пт


class ChartStyle(_Strict):
    palette: list[Hex] = Field(min_length=1)  # серии по кругу; первый — рост у гистограммы
    negative: Hex                             # столбцы с отрицательным значением


class TrendStyle(_Strict):
    up: Hex    # изменение метрики со знаком «+»
    down: Hex  # изменение со знаком «-» или тире


class TableStyle(_Strict):
    size: float = Field(gt=0)  # кегль ячеек, пт
    header_fill: Hex           # заливка шапки
    stripe_fill: Hex           # заливка чётных строк («зебра»)
    rule: Hex                  # цвет тонких линий


class SlideMap(_Strict):
    layout: str                                   # имя макета мастера (<p:cSld name>)
    # поле → idx заполнителя; у заголовка 0. Список — только у columns.blocks
    ph: dict[str, int | list[int]]
    body: tuple[float, float, float, float] | None = None  # зона тела x, y, w, h в мм
    image_width: dict[Literal["small", "medium", "large"], Percent] | None = None
    text_width: Percent | None = None
    gap: float | None = Field(default=None, ge=0)  # зазор между колонками, мм
    heading: str | None = None                     # заголовок колонки тезисов у графиков


SlideType = Literal["title", "agenda", "section", "text", "image", "columns", "comparison",
                    "quote", "chart", "dual-chart", "metrics", "table", "gallery", "video", "final"]
assert set(SlideType.__args__) == set(SLIDE_TYPES), "pptx.SlideType разошёлся с model.SLIDE_TYPES"


class Mapping(_Strict):
    master: str
    # шрифт всего текста в PPTX; None — шрифты мастера как есть (см. _set_font)
    font: str | None = Field(default=None, min_length=1)
    text: TextStyle
    chart: ChartStyle
    trend: TrendStyle
    table: TableStyle
    slides: dict[SlideType, SlideMap]


def ph_fields(stype: str) -> tuple[str, ...]:
    return PH_FIELDS.get(stype, ("title",))


def _body_keys(stype: str, sm: SlideMap) -> tuple[str, ...]:
    if stype == "columns" and "blocks" in sm.ph:
        return ()  # блоки в заполнителях мастера — зона и зазор не нужны
    return BODY_KEYS.get(stype, ())


def _completeness(m: Mapping) -> list[str]:
    """Правила полноты спеки §7, которые не выражаются типами."""
    problems = [f"slides: нет типа {t}" for t in SLIDE_TYPES if t not in m.slides]
    for stype, sm in m.slides.items():
        need = ph_fields(stype)
        allowed_ph = need + (("blocks",) if stype == "columns" else ())
        problems += [f"slides.{stype}.ph: нет поля {f}" for f in need if f not in sm.ph]
        for f, v in sm.ph.items():
            if f not in allowed_ph:
                problems.append(f"slides.{stype}.ph: поле {f} не кладётся в заполнитель")
            elif f == "blocks":
                if not isinstance(v, list) or len(v) < BLOCKS:
                    problems.append(f"slides.{stype}.ph.blocks: нужно не меньше {BLOCKS} заполнителей")
            elif isinstance(v, list):
                problems.append(f"slides.{stype}.ph.{f}: ожидается один idx, а не список")
        body = _body_keys(stype, sm)
        allowed = body + OPTIONAL_KEYS.get(stype, ())
        for key in ("body", "image_width", "text_width", "gap", "heading"):
            value = getattr(sm, key)
            if key in body and value is None:
                problems.append(f"slides.{stype}: нет {key}")
            elif key not in allowed and value is not None:
                problems.append(f"slides.{stype}: {key} у этого типа не бывает")
        if sm.image_width is not None and len(sm.image_width) != 3:
            problems.append(f"slides.{stype}.image_width: нужны small, medium и large")
        if stype == "image" and sm.body is not None and sm.image_width is not None and sm.gap is not None:
            body_w, body_h = sm.body[2], sm.body[3]
            p = max(sm.image_width.values())
            if not (body_w * (100 - p) / 100 - sm.gap > 0):
                problems.append(f"slides.image: при image_width {p:g} % и gap {sm.gap:g} мм "
                                 "колонке тезисов не остаётся места")
            if not (body_h > CAPTION_STRIP):
                problems.append(f"slides.image.body: высота {body_h:g} мм не больше "
                                 f"полосы подписи {CAPTION_STRIP} мм")
        if stype == "chart" and sm.body is not None and sm.text_width is not None and sm.gap is not None:
            body_w = sm.body[2]
            if not (body_w * (100 - sm.text_width) / 100 - sm.gap > 0):
                problems.append(f"slides.{stype}: при text_width {sm.text_width:g} % и gap {sm.gap:g} мм "
                                 "графику не остаётся места")
    return problems


def load_mapping(template_dir: Path, rel: object) -> Mapping:
    """Сопоставление из шаблона: структура, полнота и путь к мастеру; без сверки с мастером."""
    try:
        path = inside_template(template_dir, rel, "pptx")
    except ValueError as e:
        raise MappingError(str(e)) from None
    try:
        raw = json.loads(path.read_text())
    except FileNotFoundError:
        raise MappingError(f"pptx: нет файла {rel}") from None
    except (OSError, ValueError) as e:
        raise MappingError(f"{rel} не читается: {e}") from None
    try:
        m = Mapping.model_validate(raw)
    except ValidationError as e:
        raise MappingError(f"{rel}: " + "; ".join(
            f"{'.'.join(map(str, x['loc']))}: {x['msg']}" for x in e.errors())) from None
    problems = _completeness(m)
    try:
        master = inside_template(template_dir, m.master, "master")
    except ValueError as e:
        problems.append(str(e))
    else:
        if not master.is_file():
            problems.append(f"master: нет файла {m.master}")
    if problems:
        raise MappingError(f"{rel}: " + "; ".join(problems))
    return m


def _layouts(prs) -> dict:
    """Макеты всех мастеров по имени; при совпадении имён побеждает первый мастер."""
    out: dict = {}
    for master in prs.slide_masters:
        for layout in master.slide_layouts:
            out.setdefault(layout.name, layout)
    return out


# add_slide не копирует эти заполнители макета (кроме номера слайда — его
# переносит _slide_number отдельно): поле, сопоставленное на их idx,
# провалилось бы молча, поэтому open_master их из «have» не считает.
_NOT_COPIED = {PP_PLACEHOLDER.DATE, PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.SLIDE_NUMBER}


def open_master(template_dir: Path, m: Mapping):
    """Мастер шаблона, сверенный с сопоставлением: макеты, заполнители, зоны."""
    try:
        prs = Presentation(str(Path(template_dir) / m.master))
    except Exception as e:  # битый zip, не презентация — для автора шаблона это одна ошибка
        raise MappingError(f"мастер {m.master} не открывается: {type(e).__name__}: {e}") from None
    layouts = _layouts(prs)
    w, h = prs.slide_width / MM, prs.slide_height / MM
    problems = []
    for stype, sm in m.slides.items():
        layout = layouts.get(sm.layout)
        if layout is None:
            problems.append(f"slides.{stype}: в мастере нет макета «{sm.layout}»")
            continue
        have = {ph.placeholder_format.idx for ph in layout.placeholders
                if ph.placeholder_format.type not in _NOT_COPIED}
        for f, i in sm.ph.items():
            for idx in (i if isinstance(i, list) else [i]):
                if idx not in have:
                    problems.append(f"slides.{stype}.ph.{f}: в макете «{sm.layout}» нет заполнителя idx {idx}")
        if sm.body is not None:
            x, y, bw, bh = sm.body
            if bw <= 0 or bh <= 0 or x < 0 or y < 0 or x + bw > w + 0.5 or y + bh > h + 0.5:
                problems.append(f"slides.{stype}.body: зона {list(sm.body)} выходит за слайд {w:.0f}×{h:.0f} мм")
    if problems:
        raise MappingError("; ".join(problems))
    return prs


def clean(text: str) -> str:
    return CONTROL.sub("", text)


def _drop_slides(prs) -> None:
    """Слайды-образцы мастера не нужны: сначала ссылка из списка, потом связь.
    Недостижимые части python-pptx при сохранении не пишет."""
    ids = prs.slides._sldIdLst
    for sld in list(ids):
        ids.remove(sld)
        prs.part.drop_rel(sld.rId)


def _retype(root, font: str) -> None:
    """Явные латинские гарнитуры (a:latin) — на font. Ссылки на шрифт темы
    (+mj-lt, +mn-lt) остаются ссылками, слоты ea/cs/sym не трогаются: кириллицу
    и латиницу PowerPoint набирает из a:latin."""
    for latin in root.iter(qn("a:latin")):
        if not latin.get("typeface", "").startswith("+"):
            latin.set("typeface", font)


def _retype_part(part, font: str) -> None:
    """Мастера и макеты python-pptx разбирает в XML, а темы и мастер раздаточных
    материалов держит байтами — те правятся как XML и кладутся обратно."""
    if hasattr(part, "_element"):
        _retype(part._element, font)
        return
    root = etree.fromstring(part.blob)
    _retype(root, font)
    part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def _set_font(prs, font: str) -> None:
    """Шрифт шаблона для всего текста: мастера слайдов с макетами, мастер заметок и
    мастер раздаточных материалов — и тема каждого (заголовки, основной текст,
    умолчания новых фигур). Отсутствующие мастера не создаются. Зовётся после
    сборки слайдов: мастер заметок python-pptx заводит на первой заметке."""
    parts = []
    for master in prs.slide_masters:
        parts += [master.part, *(layout.part for layout in master.slide_layouts)]
    for reltype in (RT.NOTES_MASTER, RT.HANDOUT_MASTER):
        try:
            parts.append(prs.part.part_related_by(reltype))
        except KeyError:
            pass
    themes = []
    for part in parts:
        try:
            theme = part.part_related_by(RT.THEME)
        except KeyError:  # у макетов своей темы нет
            continue
        if theme not in themes:
            themes.append(theme)
    for part in parts + themes:
        _retype_part(part, font)


@dataclass
class Ctx:
    """Всё, что нужно функции тела слайда; собирается в export_pptx на каждый слайд."""
    slide: object
    s: object
    sm: SlideMap
    m: Mapping
    assets: Path
    tmp: Path
    typst_bin: str
    warnings: list[str]
    n: int


def _fill(slide, s, sm: SlideMap, keep: frozenset[int] | set[int] = frozenset()) -> None:
    """Поля — в заполнители по idx; пустые и несопоставленные заполнители удаляются,
    иначе при правке в PowerPoint видна подсказка макета («Уровень текста 1…»).
    keep — заполнители, которые уже заполнила функция тела (блоки columns)."""
    groups: dict[int, list[str]] = {}
    for f in ph_fields(s.type):
        value = clean(getattr(s, f) or "")
        if value.strip():
            groups.setdefault(sm.ph[f], []).append(value)
    for ph in list(slide.placeholders):
        idx = ph.placeholder_format.idx
        if idx in keep:
            continue
        texts = groups.get(idx)
        if not texts:
            ph._element.getparent().remove(ph._element)
            continue
        tf = ph.text_frame
        for i, t in enumerate(texts):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = t.replace("\n", "\v")  # \v — разрыв строки внутри абзаца
        if idx == 0:
            # автоподбор PowerPoint: пересчитывается, когда заголовок правят
            tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE


def _slide_number(slide, layout, n: int) -> None:
    """python-pptx номер слайда не переносит: копия заполнителя макета с полем номера,
    чтобы PowerPoint пересчитал его при перестановке слайдов."""
    for ph in layout.placeholders:
        if ph.placeholder_format.type != PP_PLACEHOLDER.SLIDE_NUMBER:
            continue
        el = copy.deepcopy(ph._element)
        el.find(qn("p:nvSpPr")).find(qn("p:cNvPr")).set("id", str(slide.shapes._next_shape_id))
        body = el.find(qn("p:txBody"))
        for p in body.findall(qn("a:p")):
            body.remove(p)
        p = etree.SubElement(body, qn("a:p"))
        fld = etree.SubElement(p, qn("a:fld"), id="{" + str(uuid.uuid4()).upper() + "}", type="slidenum")
        etree.SubElement(fld, qn("a:t")).text = str(n)
        slide.shapes._spTree.append(el)
        return


def _paragraph(tf, k: int, text: str, size: float, *, bold: bool = False, color: str | None = None,
               align=None, keep_newlines: bool = True):
    """k-й абзац надписи: первый уже есть, остальные добавляются."""
    p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
    text = clean(text)
    p.text = text.replace("\n", "\v") if keep_newlines else text.replace("\n", " ")
    if align is not None:
        p.alignment = align
    for r in p.runs:
        r.font.size = Pt(size)
        if bold:
            r.font.bold = True
        if color:
            r.font.color.rgb = RGBColor.from_string(color.upper())
    return p


def _textbox(slide, box):
    x, y, w, h = box
    tf = slide.shapes.add_textbox(Emu(x), Emu(y), Emu(w), Emu(h)).text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    return tf


def _bullets(slide, box, items: list[str], style: TextStyle, heading: str | None = None,
             numbered: bool = False) -> None:
    """Тезисы надписью в зоне: маркер (или номер) цвета style.bullet, шрифт — из темы мастера.
    Порядок в a:pPr — по схеме: spcBef, buClr, затем buAutoNum или buChar."""
    tf = _textbox(slide, box)
    paras = ([(heading, False)] if heading else []) + [(i, True) for i in items]
    indent = Mm(10 if numbered else 8)
    for k, (text, bullet) in enumerate(paras):
        p = _paragraph(tf, k, text, style.size, bold=not bullet)
        if k:
            p.space_before = Pt(style.size * 0.5)
        if bullet:
            ppr = p._p.get_or_add_pPr()
            ppr.set("marL", str(indent))
            ppr.set("indent", str(-indent))
            clr = etree.SubElement(ppr, qn("a:buClr"))
            etree.SubElement(clr, qn("a:srgbClr"), val=style.bullet.upper())
            if numbered:
                etree.SubElement(ppr, qn("a:buAutoNum"), type="arabicPeriod")
            else:
                etree.SubElement(ppr, qn("a:buChar"), char="•")


def _split(box, n: int, gap_mm: float) -> list[tuple[int, int, int, int]]:
    """n равных колонок зоны с зазором gap_mm между ними."""
    x, y, w, h = box
    g = int(gap_mm * MM)
    cw = (w - g * (n - 1)) // n
    return [(x + i * (cw + g), y, cw, h) for i in range(n)]


def _mm_box(zone) -> tuple[int, int, int, int]:
    return tuple(int(v * MM) for v in zone)


def _columns(box, left_pct: float, gap_mm: float) -> tuple[tuple, tuple]:
    """Левая колонка шириной left_pct % зоны и правая — остаток за вычетом зазора."""
    x, y, w, h = box
    lw = int(w * left_pct / 100)
    g = int(gap_mm * MM)
    return (x, y, lw, h), (x + lw + g, y, w - lw - g, h)


def _picture(slide, path: Path, box) -> int:
    """Картинка вписана в box с сохранением пропорций, по центру и к верху; возвращает её низ."""
    x, y, w, h = box
    with Image.open(path) as im:
        iw, ih = im.size
    k = min(w / iw, h / ih)
    pw, ph = int(iw * k), int(ih * k)
    slide.shapes.add_picture(str(path), Emu(x + (w - pw) // 2), Emu(y), Emu(pw), Emu(ph))
    return y + ph


def _caption(slide, text: str, x: int, y: int, w: int, style: TextStyle) -> None:
    tf = slide.shapes.add_textbox(Emu(x), Emu(y), Emu(w), Emu(int((CAPTION_STRIP - 2) * MM))).text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = clean(text).replace("\n", " ")
    for r in p.runs:
        r.font.size = Pt(style.size * CAPTION_SCALE)
        r.font.italic = True
        r.font.color.rgb = RGBColor.from_string(style.caption.upper())


def _image_file(name: str, assets_dir: Path, tmp: Path, typst_bin: str,
                warnings: list[str], n: int) -> Path | None:
    """Файл картинки; SVG — растром. Пропавшая или нерастеризуемая — None и предупреждение,
    как у PDF: слайд собирается без неё."""
    if not name:
        return None
    if Path(name).name != name:  # «/», «..»-путь или абсолютный путь — за пределы assets/
        warnings.append(f"слайд {n}: картинка {name} не найдена")
        return None
    src = Path(assets_dir) / name
    if not src.is_file():
        warnings.append(f"слайд {n}: картинка {name} не найдена")
        return None
    if src.suffix.lower() != ".svg":
        try:
            with Image.open(src) as im:
                fmt = im.format
        except OSError:  # в т.ч. PIL.UnidentifiedImageError — подтип OSError
            fmt = None
        if fmt not in {"PNG", "JPEG", "GIF", "BMP", "TIFF"}:
            warnings.append(f"слайд {n}: картинка {name} не читается")
            return None
        return src
    try:
        return rasterize_svg(src, tmp / f"{src.stem}.png", typst_bin)
    except CompileError as e:
        warnings.append(f"слайд {n}: SVG {name} не растеризуется: {e}")
        return None


def _image_body(c: Ctx) -> set[int]:
    slide, s, sm, style = c.slide, c.s, c.sm, c.m.text
    img = _image_file(s.image, c.assets, c.tmp, c.typst_bin, c.warnings, c.n)
    box = _mm_box(sm.body)
    if img is None:
        if s.bullets:
            _bullets(slide, box, s.bullets, style)
        return set()
    # картинка — ровно image_width % зоны справа, тезисы слева — остаток за вычетом зазора (спека §5.3)
    x, y, w, h = box
    iw = int(w * sm.image_width[s.image_size] / 100)
    g = int(sm.gap * MM)
    text_box, img_box = (x, y, w - iw - g, h), (x + w - iw, y, iw, h)
    if s.bullets:
        _bullets(slide, text_box, s.bullets, style)
    x, y, w, h = img_box
    strip = int(CAPTION_STRIP * MM) if s.caption else 0
    bottom = _picture(slide, img, (x, y, w, h - strip))
    if s.caption:
        _caption(slide, s.caption, x, bottom + int(2 * MM), w, style)
    return set()


def _number_format(unit: str) -> str:
    unit = unit.replace('"', "").strip()
    return f'General" {unit}"' if unit else "General"


def _bar_chart(slide, s, box, cs: ChartStyle) -> None:
    """Гистограмма: цвет столбца по знаку — рост palette[0], падение negative.

    number_format не идёт в CategoryChartData: python-pptx пишет его в
    <c:formatCode> текстовым узлом без экранирования (pptx/chart/xmlwriter.py),
    и «&», «<» в unit ломают XML графика. plot.data_labels.number_format —
    атрибут, его пишет lxml и экранирует сам (спека §5.5)."""
    unit = clean(s.unit)
    fmt = _number_format(unit)
    data = CategoryChartData()
    data.categories = [clean(d.label) for d in s.data]
    data.add_series(unit or "Значение", [d.value for d in s.data])
    x, y, w, h = box
    chart = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Emu(x), Emu(y), Emu(w), Emu(h), data).chart
    chart.has_legend = False
    chart.value_axis.has_major_gridlines = False
    plot = chart.plots[0]
    plot.has_data_labels = True
    plot.data_labels.number_format = fmt
    plot.data_labels.number_format_is_linked = False
    plot.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
    series = plot.series[0]
    series.invert_if_negative = False
    for i, d in enumerate(s.data):
        fill = series.points[i].format.fill
        fill.solid()
        fill.fore_color.rgb = RGBColor.from_string((cs.negative if d.value < 0 else cs.palette[0]).upper())


def _line_chart(slide, s, box, cs: ChartStyle) -> None:
    """Линейный график: серии по палитре по кругу, легенда сверху, единица — подпись оси."""
    data = CategoryChartData()
    data.categories = [clean(lb) for lb in s.labels]
    for ser in s.series:
        data.add_series(clean(ser.name), ser.values)
    x, y, w, h = box
    chart = slide.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, Emu(x), Emu(y), Emu(w), Emu(h), data).chart
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.TOP
    chart.legend.include_in_layout = False
    chart.value_axis.has_major_gridlines = False
    if s.unit:
        chart.value_axis.has_title = True
        chart.value_axis.axis_title.text_frame.text = clean(s.unit)
    for i, series in enumerate(chart.plots[0].series):
        color = RGBColor.from_string(cs.palette[i % len(cs.palette)].upper())
        series.smooth = False
        series.format.line.color.rgb = color
        series.format.line.width = Pt(3)
        series.marker.format.fill.solid()
        series.marker.format.fill.fore_color.rgb = color
        series.marker.format.line.color.rgb = color


def _chart_body(c: Ctx) -> set[int]:
    box = _mm_box(c.sm.body)
    if not c.s.bullets:
        _line_chart(c.slide, c.s, box, c.m.chart)
        return set()
    text_box, chart_box = _columns(box, c.sm.text_width, c.sm.gap)
    _bullets(c.slide, text_box, c.s.bullets, c.m.text, heading=c.sm.heading)
    _line_chart(c.slide, c.s, chart_box, c.m.chart)
    return set()


def _agenda_body(c: Ctx) -> set[int]:
    _bullets(c.slide, _mm_box(c.sm.body), c.s.items, c.m.text, numbered=True)
    return set()


def _text_body(c: Ctx) -> set[int]:
    _bullets(c.slide, _mm_box(c.sm.body), c.s.bullets, c.m.text)
    return set()


def _block_text(tf, block, style: TextStyle) -> None:
    """Блок columns: жирный заголовок (если есть), затем текст."""
    k = 0
    if clean(block.heading).strip():
        _paragraph(tf, k, block.heading, style.size, bold=True)
        k += 1
    p = _paragraph(tf, k, block.text, style.size)
    if k:
        p.space_before = Pt(style.size * 0.5)


def _columns_body(c: Ctx) -> set[int]:
    idxs = c.sm.ph.get("blocks")
    if idxs:  # блоки в заполнителях мастера; лишние удалит _fill — их нет в keep
        by_idx = {ph.placeholder_format.idx: ph for ph in c.slide.placeholders}
        used = set()
        for block, idx in zip(c.s.blocks, idxs):
            tf = by_idx[idx].text_frame
            tf.clear()
            _block_text(tf, block, c.m.text)
            used.add(idx)
        return used
    for block, box in zip(c.s.blocks, _split(_mm_box(c.sm.body), len(c.s.blocks), c.sm.gap)):
        _block_text(_textbox(c.slide, box), block, c.m.text)
    return set()


def _comparison_body(c: Ctx) -> set[int]:
    left, right = _split(_mm_box(c.sm.body), 2, c.sm.gap)
    for side, box in ((c.s.left, left), (c.s.right, right)):
        _bullets(c.slide, box, side.bullets, c.m.text, heading=clean(side.heading).strip() or None)
    return set()


def _quote_body(c: Ctx) -> set[int]:
    tf = _textbox(c.slide, _mm_box(c.sm.body))
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    _paragraph(tf, 0, c.s.text, c.m.text.quote, align=PP_ALIGN.CENTER)
    if clean(c.s.author).strip():
        p = _paragraph(tf, 1, c.s.author, c.m.text.size, color=c.m.text.caption, align=PP_ALIGN.CENTER)
        p.space_before = Pt(c.m.text.size)
    return set()


BODIES: dict[str, Callable[[Ctx], set[int]]] = {
    "agenda": _agenda_body, "text": _text_body, "image": _image_body, "columns": _columns_body,
    "comparison": _comparison_body, "quote": _quote_body, "chart": _chart_body,
}


def export_pptx(deck: Deck, template_dir: Path, mapping_rel: object, assets_dir: Path,
                typst_bin: str = "typst", warnings: list[str] | None = None) -> bytes:
    """Колода → байты .pptx на мастере шаблона.

    Сопоставление читается на каждый вызов: канон могли обновить после старта,
    и ошибка должна прийти MappingError с причиной, а не 500. Пропавшие картинки
    и нерастеризуемые SVG — в warnings, слайд собирается без них.
    """
    warnings = [] if warnings is None else warnings
    m = load_mapping(template_dir, mapping_rel)
    prs = open_master(template_dir, m)
    layouts = _layouts(prs)
    _drop_slides(prs)
    with tempfile.TemporaryDirectory() as tmp:
        for n, s in enumerate(deck.slides, 1):
            sm = m.slides[s.type]
            layout = layouts[sm.layout]
            slide = prs.slides.add_slide(layout)
            body = BODIES.get(s.type)
            keep = body(Ctx(slide, s, sm, m, Path(assets_dir), Path(tmp), typst_bin, warnings, n)) if body else set()
            _fill(slide, s, sm, keep)
            _slide_number(slide, layout, n)
            if s.notes.strip():
                slide.notes_slide.notes_text_frame.text = clean(s.notes)
        if m.font:
            _set_font(prs, m.font)
        buf = io.BytesIO()
        prs.save(buf)
    return buf.getvalue()
