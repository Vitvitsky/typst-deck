"""Текст из загруженных файлов для контекста LLM."""
from __future__ import annotations

import csv
import io
from dataclasses import dataclass
from pathlib import Path
from typing import Literal

IMAGE_EXT = {".png", ".jpg", ".jpeg", ".svg"}
MAX_ROWS = 200


class UnsupportedFile(ValueError):
    pass


@dataclass
class Extracted:
    kind: Literal["text", "image"]
    text: str = ""
    truncated: bool = False


def _md_table(rows: list[list]) -> str:
    rows = [[str(c) if c is not None else "" for c in r] for r in rows]
    if not rows:
        return ""
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    lines = ["| " + " | ".join(rows[0]) + " |", "|" + "---|" * width]
    lines += ["| " + " | ".join(r) + " |" for r in rows[1:]]
    return "\n".join(lines)


def _docx(data: bytes) -> str:
    import docx

    d = docx.Document(io.BytesIO(data))
    parts = [p.text for p in d.paragraphs if p.text.strip()]
    for t in d.tables:
        parts.append(_md_table([[c.text for c in r.cells] for r in t.rows]))
    return "\n\n".join(parts)


def _xlsx(data: bytes) -> str:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    parts = []
    for ws in wb.worksheets:
        rows: list[list] = []
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= MAX_ROWS:
                rows.append(["…"])
                break
            rows.append(list(row))
        parts.append(f"## {ws.title}\n\n{_md_table(rows)}")
    return "\n\n".join(parts)


def _csv(data: bytes) -> str:
    text = data.decode("utf-8", errors="replace")
    return _md_table(list(csv.reader(io.StringIO(text)))[: MAX_ROWS + 1])


def extract(name: str, data: bytes, limit: int = 60_000) -> Extracted:
    """Картинки распознаются по расширению, остальное превращается в текст."""
    ext = Path(name).suffix.lower()
    if ext in IMAGE_EXT:
        return Extracted("image")
    if ext in (".txt", ".md"):
        text = data.decode("utf-8", errors="replace")
    elif ext == ".docx":
        text = _docx(data)
    elif ext == ".xlsx":
        text = _xlsx(data)
    elif ext == ".csv":
        text = _csv(data)
    else:
        raise UnsupportedFile(f"формат {ext or '(без расширения)'} не поддерживается")
    truncated = len(text) > limit
    if truncated:
        text = text[:limit] + "\n\n[… усечено: файл длиннее лимита]"
    return Extracted("text", text, truncated)
